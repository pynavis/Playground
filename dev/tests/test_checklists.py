# -*- coding: utf-8 -*-
"""namo.checklists: templates, records, progress, sign-off, merging and the reports."""
import datetime
import os

import _compat
from namo import checklists as cl

D = datetime.datetime


def test_template_text_round_trip_and_builtins():
    items = cl.parse_template_text('# Setting out\nGridlines checked\n\nLevels checked\n# Fire\n  Stopping done  ')
    assert items == [{'id': 'I01', 'section': 'Setting out', 'text': 'Gridlines checked'},
                     {'id': 'I02', 'section': 'Setting out', 'text': 'Levels checked'},
                     {'id': 'I03', 'section': 'Fire', 'text': 'Stopping done'}]
    assert cl.parse_template_text(cl.template_text(items)) == items
    _compat.raises(ValueError, cl.parse_template_text, '# Only a section\n\n')
    for name, config in cl.BUILTIN.items():
        assert len(cl.template_items(config)) >= 8, name
    merged = cl.all_templates({'QA inspection': {'text': 'One thing'}, 'Mine': {'text': 'A'}})
    assert sorted(merged) == ['Mine', 'QA inspection', 'Safety walk']
    assert cl.template_items(merged['QA inspection'])[0]['text'] == 'One thing'
    assert cl.template_items({'items': [{'text': 'x'}]}) == [{'id': 'I01', 'section': '', 'text': 'x'}]


def _record():
    return cl.new_checklist('{guid-1}', 'Core L2', 'QA/Level 2', 'Mini', {'text': '# A\nOne\nTwo\n# B\nThree'},
                            now=D(2026, 10, 5, 8, 0, 0))


def test_progress_results_and_sign_off():
    record = _record()
    assert record['id'] == '{guid-1}' and cl.path_of(record) == 'QA/Level 2/Core L2'
    assert cl.progress(record) == {'total': 3, 'answered': 0, 'passed': 0, 'failed': 0, 'na': 0,
                                   'open': 3, 'percent': 0, 'complete': False, 'signed': False}
    cl.set_result(record, 'I01', 'pass')
    cl.set_result(record, 'I02', 'fail', 'Spacing 300 not 200')
    assert cl.progress_text(record) == '2 of 3 answered, 1 failed'
    _compat.raises(ValueError, cl.sign_off, record, 'Ana')
    cl.set_result(record, 'I03', 'na')
    cl.set_result(record, 'I03', note='Not built yet')
    assert record['items'][2] == {'id': 'I03', 'section': 'B', 'text': 'Three', 'result': 'na',
                                  'note': 'Not built yet'}
    _compat.raises(ValueError, cl.sign_off, record, '  ')
    cl.sign_off(record, 'Ana', now=D(2026, 10, 6, 9, 0, 0))
    assert cl.progress(record)['signed'] and record['signed_at'] == '2026-10-06T09:00:00'
    assert cl.progress_text(record) == '3 of 3 answered, 1 failed, signed off'
    cl.sign_off(record, '', on=False)
    assert not record['signed_off'] and record['signed_by'] == ''
    cl.set_result(record, 'I01', None)
    assert record['items'][0]['result'] is None
    _compat.raises(ValueError, cl.set_result, record, 'I01', 'maybe')
    _compat.raises(KeyError, cl.set_result, record, 'I99', 'pass')


def test_merge_template_keeps_answers():
    record = _record()
    cl.set_result(record, 'I01', 'pass')
    added = cl.merge_template(record, {'text': '# A\nOne\nTwo\nFour\n# B\nThree\nFive'})
    assert added == 2
    assert [i['text'] for i in record['items']] == ['One', 'Two', 'Three', 'Four', 'Five']
    assert record['items'][0]['result'] == 'pass' and record['items'][3]['id'] == 'I04'


def _records():
    a = _record()
    for item_id, result in (('I01', 'pass'), ('I02', 'fail'), ('I03', 'pass')):
        cl.set_result(a, item_id, result, 'note %s' % item_id)
    b = cl.new_checklist('{guid-2}', 'Roof', '', 'Mini', {'text': 'One\nTwo'})
    cl.set_result(b, 'I01', 'pass')
    return [a, b]


def test_summary_failures_totals():
    records = _records()
    assert cl.summary_rows(records) == [
        ['QA/Level 2/Core L2', 'Mini', '3/3', 100, 1, '', '', 'No'],
        ['Roof', 'Mini', '1/2', 50, 0, '', '', 'No']]
    assert cl.failures(records) == [['QA/Level 2/Core L2', 'A', 'Two', 'note I02']]
    assert cl.totals(records) == {'checklists': 2, 'items': 5, 'answered': 4, 'failed': 1, 'signed': 0,
                                  'percent': 80}


def test_reports():
    records = _records()
    html = cl.build_report(records, 'Field checklists', 'Tower.nwf', 'amol').to_string()
    assert 'Field checklists' in html and '80%' in html and 'note I02' in html and 'Core L2' in html
    path = os.path.join(_compat.tmpdir('checklists'), 'checklists.xlsx')
    cl.write_xlsx(path, records, 'Field checklists')
    import zipfile
    with zipfile.ZipFile(path) as package:
        assert len([n for n in package.namelist() if n.startswith('xl/worksheets/sheet')]) == 3
    if _compat.IRONPYTHON:
        return
    import openpyxl
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ['Summary', 'Failures', 'Items']
    assert [c.value for c in book['Failures'][2]] == ['QA/Level 2/Core L2', 'A', 'Two', 'note I02']
    assert book['Items'].max_row == 6
