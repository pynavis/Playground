# -*- coding: utf-8 -*-
"""The four G9 report modules through namo.reports.run, with the Navisworks reads patched
out: the rep_<id>.py contract (defaults, describe, result dict, files written)."""
import contextlib
import datetime
import os
import zipfile

import _compat
from namo import fourd_common, items, reports, takeoff, timeliner

D = datetime.datetime


@contextlib.contextmanager
def patched(module, **values):
    saved = dict((name, getattr(module, name)) for name in values)
    for name, value in values.items():
        setattr(module, name, value)
    try:
        yield
    finally:
        for name, value in saved.items():
            setattr(module, name, value)


class Coll(list):
    @property
    def Count(self):
        return len(self)


def _rows():
    return [
        {'key': (0,), 'parent_key': None, 'depth': 0, 'outline': '1', 'name': 'Level 1', 'id': '1',
         'start': D(2026, 1, 1), 'end': D(2026, 2, 1), 'actual_start': None, 'actual_end': None,
         'type': '', 'progress': None, 'enabled': True, 'user': ['Zone A'] + [''] * 9,
         'attach_kind': 'none', 'attach_text': '', 'has_children': True, 'task': 't0'},
        {'key': (0, 0), 'parent_key': (0,), 'depth': 1, 'outline': '1.1', 'name': 'Columns', 'id': '1.1',
         'start': D(2026, 1, 1), 'end': D(2026, 1, 10), 'actual_start': D(2026, 1, 2), 'actual_end': None,
         'type': 'Construct', 'progress': 40.0, 'enabled': True, 'user': [''] * 10,
         'attach_kind': 'sets', 'attach_text': 'Columns', 'has_children': False, 'task': 't1'},
        {'key': (0, 1), 'parent_key': (0,), 'depth': 1, 'outline': '1.2', 'name': 'Slab', 'id': '1.2',
         'start': D(2026, 1, 10), 'end': D(2026, 1, 5), 'actual_start': None, 'actual_end': None,
         'type': 'Construct', 'progress': None, 'enabled': True, 'user': [''] * 10,
         'attach_kind': 'items', 'attach_text': '', 'has_children': False, 'task': 't2'},
        {'key': (0, 2), 'parent_key': (0,), 'depth': 1, 'outline': '1.3', 'name': 'Walls', 'id': '1.3',
         'start': None, 'end': None, 'actual_start': None, 'actual_end': None,
         'type': 'Construct', 'progress': None, 'enabled': True, 'user': [''] * 10,
         'attach_kind': 'none', 'attach_text': '', 'has_children': False, 'task': 't3'},
    ]


def _sheets(path):
    with zipfile.ZipFile(path) as package:
        text = package.read('xl/workbook.xml').decode('utf-8')
    return [part.split('"')[0] for part in text.split('<sheet name="')[1:]]


def test_registry_finds_the_modules():
    for report_id in ('timeliner_export', 'schedule_health', 'quick_takeoff', 'zone_takeoff'):
        assert reports.module_exists(report_id)
        module = reports.module_for(report_id)
        config = module.default_config()
        assert 'folder' in config and 'name_template' in config
        assert module.describe(config)
        assert module.TITLE == reports.title_of(report_id)


def test_timeliner_export_writes_the_workbook():
    folder = _compat.tmpdir('rep_timeliner_export')
    selected = {'t1': Coll(['a', 'b', 'c']), 't2': Coll(['d'])}
    with patched(timeliner, snapshot=lambda doc=None: _rows(),
                 selected_items=lambda task, doc=None: selected.get(task)):
        with patched(fourd_common, doc_title=lambda doc=None: 'Tower.nwf'):
            result = reports.run('timeliner_export', {'folder': folder, 'name_template': 'tasks'})
    assert result['level'] == 'success', result
    assert result['message'] == 'Exported 4 tasks' and result['paths'] == [os.path.join(folder, 'tasks.xlsx')]
    assert _sheets(result['paths'][0]) == ['Tasks', 'Summary']
    if not _compat.IRONPYTHON:
        import openpyxl
        sheet = openpyxl.load_workbook(result['paths'][0])['Tasks']
        values = [[c.value for c in row] for row in sheet.iter_rows(min_row=3)]
        assert values[0][:3] == ['Task ID', 'WBS', 'Name'] and values[0][-1] == 'User1'
        assert values[2][:3] == ['1.1', '1.1', 'Columns'] and values[2][13] == 3
        assert sheet.cell(row=5, column=3).alignment.indent == 1
    with patched(timeliner, snapshot=lambda doc=None: []):
        empty = reports.run('timeliner_export', {'folder': folder})
    assert empty['level'] == 'info' and empty['message'] == 'No TimeLiner tasks'


def test_schedule_health_report():
    folder = _compat.tmpdir('rep_schedule_health')
    attached = {'t1': Coll(['col1', 'col2']), 't2': Coll(['col2', 'slab'])}
    with patched(timeliner, snapshot=lambda doc=None: _rows(),
                 selected_items=lambda task, doc=None: attached.get(task)):
        with patched(items, objects=lambda roots, tick=None, budget=None: list(roots),
                     roots=lambda doc=None: ['col1', 'col2', 'slab', 'wall', 'door']):
            with patched(fourd_common, doc_title=lambda doc=None: 'Tower.nwf'):
                result = reports.run('schedule_health', {'folder': folder, 'name_template': 'health',
                                                         'create_set': False})
    assert result['level'] == 'warning', result
    assert result['message'].startswith('Schedule health ') and result['message'].endswith(' of 100')
    assert '2 tasks with problems' in result['detail'] and '2 objects in no task' in result['detail']
    path = result['paths'][0]
    assert path.endswith('health.html')
    with open(path, 'rb') as handle:
        html = handle.read().decode('utf-8')
    assert 'Schedule health' in html and 'Walls' in html and 'In more than one task' in html


def _takeoff_result():
    records = [{'name': 'W1', 'groups': ['Walls'], 'zone': 'Level 1', 'measures': {'volume': 2.0},
                'sources': {'volume': 'property'}, 'item': None},
               {'name': 'D1', 'groups': ['Ducts'], 'zone': 'Level 2', 'measures': {'volume': None},
                'sources': {}, 'item': None}]
    measures = ['count', 'volume']
    data = {'records': records, 'pairs': {'volume': [('Element', 'Volume')]}, 'sampled': False, 'total': 2,
            'geometric': 0, 'geometry_skipped': 0, 'doc_units': 'Meters', 'scope': 'model',
            'groupers': [{'special': '@category'}], 'measures': measures}
    unit_set = takeoff.display_units('m', 'Meters')
    return {'data': data, 'rows': takeoff.aggregate(records, 1, measures), 'unit_set': unit_set,
            'measures': measures, 'group_titles': ['Category'], 'notes': takeoff.notes_for(data, {}, unit_set),
            'status': takeoff.status_text(data, 2)}


def test_quick_and_zone_takeoff_reports():
    folder = _compat.tmpdir('rep_takeoffs')
    canned = _takeoff_result()
    with patched(takeoff, run=lambda config, doc=None, report=None: canned):
        with patched(fourd_common, has_models=lambda doc=None: True, doc_title=lambda doc=None: 'Tower.nwf'):
            quick = reports.run('quick_takeoff', {'folder': folder, 'name_template': 'quick'})
            zone = reports.run('zone_takeoff', {'folder': folder, 'name_template': 'zone', 'pivot_measure': 'weight'})
        with patched(fourd_common, has_models=lambda doc=None: False):
            closed = reports.run('quick_takeoff', {'folder': folder})
    assert quick['level'] == 'success' and quick['message'] == 'Takeoff of 2 objects', quick
    assert _sheets(quick['paths'][0]) == ['Takeoff', 'Objects', 'Notes']
    assert zone['level'] == 'success' and zone['message'] == 'Takeoff of 2 objects in 2 zones', zone
    assert _sheets(zone['paths'][0]) == ['Pivot', 'Zones', 'Level 1', 'Level 2', 'Notes']
    assert closed['level'] == 'info' and closed['message'] == 'No model is open'

    def no_zones(config, doc=None, report=None):
        raise ValueError('The model has no grid levels. Pick another zone source.')
    with patched(takeoff, run=no_zones):
        with patched(fourd_common, has_models=lambda doc=None: True):
            missing = reports.run('zone_takeoff', {'folder': folder})
    assert missing['level'] == 'info' and missing['detail'].startswith('The model has no grid levels')
