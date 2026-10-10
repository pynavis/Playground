# -*- coding: utf-8 -*-
"""namo.preflight: profiles, every check's verdict, summaries, and the API half through
the fake Navisworks layer; rep_preflight's files."""
import datetime
import os

import _compat

from namo import paths, preflight


def _facts(**over):
    facts = {
        'doc_file': r'C:\Jobs\Hospital.nwf', 'modified': False,
        'models': [{'name': 'ARCH.nwc', 'units': 'Millimeters'},
                   {'name': 'MEP-L1.nwc', 'units': 'Millimeters'}],
        'stale': [{'name': 'ARCH.nwc', 'status': 'Current'}, {'name': 'MEP-L1.nwc', 'status': 'Current'}],
        'sets_found': {},
        'clash': {'available': True, 'tests': [
            {'name': 'Ducts vs Pipes', 'last_run': datetime.datetime(2026, 10, 4, 9, 0)}]},
        'newest_model_time': datetime.datetime(2026, 10, 3, 12, 0),
        'new_clashes': {'available': True, 'count': 0, 'checked': 120, 'sampled': False},
        'viewpoints': 12,
        'hidden': {'found': False, 'example': None, 'checked': 5000, 'sampled': False},
        'overrides': {'readable': True, 'found': 0, 'checked': 400, 'sampled': False},
        'section_on': False,
        'timeliner': {'available': True, 'tasks': 0, 'unattached': 0, 'examples': []},
    }
    facts.update(over)
    return facts


def _all_on(**over):
    profile = preflight.normalize_profile(preflight.DEFAULT_PROFILE)
    for check in preflight.CHECK_IDS:
        profile['checks'][check]['enabled'] = True
    profile.update(over)
    return profile


def _by_id(checks):
    return dict((c['id'], c) for c in checks)


def test_default_profile_shape():
    profile = preflight.normalize_profile(preflight.DEFAULT_PROFILE)
    assert profile['name'] == 'Standard'
    assert sorted(profile['checks']) == sorted(preflight.CHECK_IDS)
    assert 'timeliner' not in preflight.enabled_checks(profile)
    assert preflight.blocking_checks(profile) == ['saved', 'expected_models', 'sets']
    assert profile['max_new_clashes'] == 0 and profile['min_viewpoints'] == 1
    assert preflight.describe_profile(profile) == '11 checks, 3 block publishing'


def test_normalize_profile_cleans_input():
    raw = {'name': '  QA  ', 'checks': {'saved': {'enabled': False}, 'units': {'level': 'FAIL'},
                                         'bogus': {'enabled': True}, 'hidden': 'yes'},
           'expected_models': 'ARCH*\n\n  MEP.nwc \nARCH*', 'required_sets': ['A/B', '', None],
           'max_new_clashes': '-7', 'min_viewpoints': 'x', 'hidden_budget': 5}
    profile = preflight.normalize_profile(raw)
    assert profile['name'] == 'QA'
    assert profile['checks']['saved']['enabled'] is False
    assert profile['checks']['units']['level'] == 'fail'
    assert 'bogus' not in profile['checks']
    assert profile['checks']['hidden'] == {'enabled': True, 'level': 'warn'}
    assert profile['expected_models'] == ['ARCH*', 'MEP.nwc']
    assert profile['required_sets'] == ['A/B']
    assert profile['max_new_clashes'] == -1
    assert profile['min_viewpoints'] == 1
    assert profile['hidden_budget'] == 1000
    assert preflight.normalize_profile(None)['name'] == 'Untitled'


def test_profile_from_choices():
    profile = preflight.profile_from_choices('Coord', ['saved', 'units', 'sets'], ['sets', 'hidden'],
                                             ['ARCH'], ['QA/Hidden'], 5, 3)
    assert preflight.enabled_checks(profile) == ['saved', 'units', 'sets']
    assert preflight.blocking_checks(profile) == ['sets']
    assert profile['checks']['hidden'] == {'enabled': False, 'level': 'fail'}
    assert profile['expected_models'] == ['ARCH'] and profile['max_new_clashes'] == 5
    assert preflight.describe_profile(profile) == '3 checks, 1 blocks publishing'


def test_pattern_matching():
    assert preflight.pattern_matches('ARCH', 'ARCH.nwc')
    assert preflight.pattern_matches('arch*', 'ARCH-L1.nwc')
    assert preflight.pattern_matches('MEP-L?.nwc', 'mep-l2.NWC')
    assert not preflight.pattern_matches('ARCH', 'ARCH-L1.nwc')
    assert not preflight.pattern_matches('', 'ARCH.nwc')
    assert preflight.missing_patterns(['ARCH*', 'STR*', 'MEP'], ['ARCH.nwc', 'mep.nwc']) == ['STR*']


def test_all_pass():
    checks = preflight.evaluate(_facts(), _all_on())
    levels = dict((c['id'], c['level']) for c in checks)
    assert [c['id'] for c in checks] == preflight.CHECK_IDS
    assert levels['saved'] == 'pass' and levels['units'] == 'pass' and levels['stale'] == 'pass'
    assert levels['expected_models'] == 'info' and levels['sets'] == 'info'
    assert levels['clash_run'] == 'pass' and levels['new_clashes'] == 'pass'
    assert levels['timeliner'] == 'info'
    assert preflight.toast_text(checks)[0] == 'success'
    assert preflight.toast_text(checks)[1] == 'Preflight passed: 9 checks'
    assert preflight.summary_text(checks) == 'passed'


def test_disabled_checks_are_left_out():
    profile = _all_on()
    profile['checks']['section']['enabled'] = False
    ids = [c['id'] for c in preflight.evaluate(_facts(), profile)]
    assert 'section' not in ids and len(ids) == len(preflight.CHECK_IDS) - 1


def test_saved_check():
    checks = _by_id(preflight.evaluate(_facts(doc_file=''), _all_on()))
    assert checks['saved']['level'] == 'fail' and 'never been saved' in checks['saved']['detail']
    checks = _by_id(preflight.evaluate(_facts(modified=True), _all_on()))
    assert checks['saved']['level'] == 'warn' and 'Hospital.nwf has unsaved changes' in checks['saved']['detail']
    profile = _all_on()
    profile['checks']['saved']['level'] = 'warn'
    assert _by_id(preflight.evaluate(_facts(doc_file=''), profile))['saved']['level'] == 'warn'


def test_units_check():
    models = [{'name': 'A.nwc', 'units': 'Millimeters'}, {'name': 'B.nwc', 'units': 'Millimeters'},
              {'name': 'C.nwc', 'units': 'Feet'}]
    check = _by_id(preflight.evaluate(_facts(models=models), _all_on()))['units']
    assert check['level'] == 'warn'
    assert check['detail'] == 'Mixed units: 2 in Millimeters, 1 in Feet (C.nwc).'
    check = _by_id(preflight.evaluate(_facts(models=[]), _all_on()))['units']
    assert check['level'] == 'info'
    check = _by_id(preflight.evaluate(_facts(models=models[:1]), _all_on()))['units']
    assert check['detail'] == 'The model is in Millimeters.'


def test_expected_models_check():
    profile = _all_on(expected_models=['ARCH', 'STR*', 'MEP-L?.nwc'])
    check = _by_id(preflight.evaluate(_facts(), profile))['expected_models']
    assert check['level'] == 'fail' and check['detail'] == 'Missing: STR*.'
    profile = _all_on(expected_models=['ARCH', 'MEP*'])
    check = _by_id(preflight.evaluate(_facts(), profile))['expected_models']
    assert check['level'] == 'pass' and 'All 2' in check['detail']


def test_stale_check():
    stale = [{'name': 'ARCH.nwc', 'status': 'Out of date'}, {'name': 'MEP.nwc', 'status': 'File not found'},
             {'name': 'STR.nwc', 'status': 'Source not found'}]
    check = _by_id(preflight.evaluate(_facts(stale=stale), _all_on()))['stale']
    assert check['level'] == 'warn'
    assert check['detail'].startswith('1 model out of date (ARCH.nwc); 1 file not found (MEP.nwc)')
    stale = [{'name': 'ARCH.nwc', 'status': 'Current'}, {'name': 'STR.nwc', 'status': 'Source not found'}]
    check = _by_id(preflight.evaluate(_facts(stale=stale), _all_on()))['stale']
    assert check['level'] == 'pass' and '1 have no source file' in check['detail']


def test_sets_check():
    profile = _all_on(required_sets=['QA/Hidden', 'Clash/Ducts'])
    check = _by_id(preflight.evaluate(_facts(sets_found={'QA/Hidden': True}), profile))['sets']
    assert check['level'] == 'fail' and check['detail'] == 'Missing: Clash/Ducts.'
    found = {'QA/Hidden': True, 'Clash/Ducts': True}
    assert _by_id(preflight.evaluate(_facts(sets_found=found), profile))['sets']['level'] == 'pass'


def test_clash_run_check():
    tests = [{'name': 'A', 'last_run': None},
             {'name': 'B', 'last_run': datetime.datetime(2026, 10, 1)},
             {'name': 'C', 'last_run': datetime.datetime(2026, 10, 5)}]
    check = _by_id(preflight.evaluate(_facts(clash={'available': True, 'tests': tests}), _all_on()))['clash_run']
    assert check['level'] == 'warn'
    assert '1 test never run (A)' in check['detail']
    assert '1 test run before the newest model data of 2026-10-03 12:00 (B)' in check['detail']
    check = _by_id(preflight.evaluate(_facts(clash={'available': True, 'tests': []}), _all_on()))['clash_run']
    assert check['level'] == 'warn' and 'no clash tests' in check['detail']
    check = _by_id(preflight.evaluate(_facts(clash={'available': False}), _all_on()))['clash_run']
    assert check['level'] == 'info'
    same = [{'name': 'A', 'last_run': datetime.datetime(2026, 10, 3, 11, 59, 59, 500000)}]
    check = _by_id(preflight.evaluate(_facts(clash={'available': True, 'tests': same}), _all_on()))['clash_run']
    assert check['level'] == 'pass'
    check = _by_id(preflight.evaluate(_facts(newest_model_time=None), _all_on()))['clash_run']
    assert check['level'] == 'pass' and 'unknown' in check['detail']


def test_new_clashes_check():
    data = {'available': True, 'count': 37, 'checked': 500, 'sampled': False}
    check = _by_id(preflight.evaluate(_facts(new_clashes=data), _all_on()))['new_clashes']
    assert check['level'] == 'warn' and check['detail'].startswith('37 New clashes; review')
    check = _by_id(preflight.evaluate(_facts(new_clashes=data), _all_on(max_new_clashes=50)))['new_clashes']
    assert check['level'] == 'pass' and 'within the limit of 50' in check['detail']
    check = _by_id(preflight.evaluate(_facts(new_clashes=data), _all_on(max_new_clashes=10)))['new_clashes']
    assert check['detail'] == '37 New clashes; the limit is 10.'
    data = {'available': True, 'count': 1, 'checked': 200000, 'sampled': True}
    check = _by_id(preflight.evaluate(_facts(new_clashes=data), _all_on()))['new_clashes']
    assert 'At least 1 New clash in the first 200,000 results' in check['detail']
    check = _by_id(preflight.evaluate(_facts(), _all_on(max_new_clashes=-1)))['new_clashes']
    assert check['level'] == 'info'


def test_viewpoints_hidden_overrides_section():
    checks = _by_id(preflight.evaluate(_facts(viewpoints=0), _all_on()))
    assert checks['viewpoints']['level'] == 'warn' and checks['viewpoints']['detail'] == 'There are no saved viewpoints.'
    checks = _by_id(preflight.evaluate(_facts(viewpoints=3), _all_on(min_viewpoints=5)))
    assert '3 saved viewpoints; this profile asks for at least 5' in checks['viewpoints']['detail']
    hidden = {'found': True, 'example': 'MEP.nwc > Hidden box', 'checked': 10, 'sampled': False}
    checks = _by_id(preflight.evaluate(_facts(hidden=hidden), _all_on()))
    assert checks['hidden']['level'] == 'warn' and 'MEP.nwc > Hidden box' in checks['hidden']['detail']
    hidden = {'found': False, 'checked': 500000, 'sampled': True}
    checks = _by_id(preflight.evaluate(_facts(hidden=hidden), _all_on()))
    assert checks['hidden']['level'] == 'pass' and 'first 500,000' in checks['hidden']['detail']
    overrides = {'readable': True, 'found': 120, 'checked': 20000, 'sampled': True}
    checks = _by_id(preflight.evaluate(_facts(overrides=overrides), _all_on()))
    assert checks['overrides']['level'] == 'warn'
    assert '120 of 20,000 geometry items sampled (best effort)' in checks['overrides']['detail']
    checks = _by_id(preflight.evaluate(_facts(overrides={'readable': False}), _all_on()))
    assert checks['overrides']['level'] == 'info'
    checks = _by_id(preflight.evaluate(_facts(section_on=True), _all_on()))
    assert checks['section']['level'] == 'warn'


def test_timeliner_check():
    data = {'available': True, 'tasks': 80, 'unattached': 12, 'examples': ['Pour L1', 'Pour L2', 'Walls']}
    check = _by_id(preflight.evaluate(_facts(timeliner=data), _all_on()))['timeliner']
    assert check['level'] == 'warn'
    assert check['detail'] == '12 of 80 tasks have nothing attached: Pour L1, Pour L2 and Walls.'
    data = {'available': True, 'tasks': 4, 'unattached': 0}
    assert _by_id(preflight.evaluate(_facts(timeliner=data), _all_on()))['timeliner']['level'] == 'pass'


def test_missing_and_broken_facts_are_info():
    checks = preflight.evaluate({}, _all_on())
    assert all(c['level'] == 'info' for c in checks)
    assert preflight.toast_text(checks)[0] == 'info'
    checks = _by_id(preflight.evaluate(_facts(models=[{'units': object()}, None]), _all_on()))
    assert checks['units']['level'] == 'info' and checks['units']['detail'].startswith('Could not be checked')


def test_summaries():
    checks = [{'id': 'saved', 'name': 'Document saved', 'level': 'fail', 'detail': ''},
              {'id': 'sets', 'name': 'Required selection sets', 'level': 'fail', 'detail': ''},
              {'id': 'units', 'name': 'Units consistent', 'level': 'warn', 'detail': ''},
              {'id': 'section', 'name': 'Sectioning off', 'level': 'pass', 'detail': ''}]
    s = preflight.summarize(checks)
    assert (s['passed'], s['warnings'], s['failures'], s['level']) == (1, 1, 2, 'fail')
    level, message, detail = preflight.toast_text(checks)
    assert level == 'error' and message == 'Preflight: 2 checks failed'
    assert detail == 'Document saved and Required selection sets; 1 warning.'
    assert preflight.summary_text(checks) == '2 failed, 1 warning'
    assert preflight.failures_text(checks) == 'Document saved and Required selection sets'
    level, message, _detail = preflight.toast_text(checks[2:])
    assert level == 'warning' and message == 'Preflight: 1 warning'
    assert preflight.toast_text([])[0] == 'info'


def test_output_html_escapes_and_has_chips():
    checks = [{'id': 'saved', 'name': 'Document saved', 'level': 'fail', 'detail': 'A <b> & c'}]
    html = preflight.output_html(checks, 'Standard', 'Job <1>.nwf')
    assert 'A &lt;b&gt; &amp; c' in html and 'Job &lt;1&gt;.nwf' in html
    assert 'var(--err)' in html and 'Failures' in html


def test_profiles_in_templates():
    folder = _compat.tmpdir('preflight_templates' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        assert preflight.profile_names() == ['Standard']
        assert preflight.load_profile('standard')['name'] == 'Standard'
        assert preflight.load_profile('Missing') is None
        custom = preflight.profile_from_choices('Coordination', ['saved'], ['saved'], [], [], 0, 0)
        assert preflight.save_profile(custom) == 'Coordination'
        own_standard = preflight.profile_from_choices('Standard', ['units'], [], [], [], 0, 0)
        preflight.save_profile(own_standard)
        assert preflight.profile_names() == ['Standard', 'Coordination']
        assert preflight.enabled_checks(preflight.load_profile('Standard')) == ['units']
        assert preflight.is_saved('Standard')
        assert preflight.delete_profile('Standard')
        assert 'timeliner' not in preflight.enabled_checks(preflight.load_profile('Standard'))
        assert preflight.load_profile('coordination')['name'] == 'coordination'
    finally:
        paths.redirect(previous)


# ---- API half through the fake layer -------------------------------------------------------


def _fake_doc():
    import test_items as nav
    doc, named = nav.model_tree()
    doc.IsModified = True
    doc.FileName = r'C:\Jobs\Hospital.nwf'
    doc.CurrentViewpoint.vp.ClipPlanes.Enabled = True
    view = nav.SavedViewpoint()
    view.DisplayName = 'Level 1'
    doc.SavedViewpoints.RootItem.Children.append(view)
    return nav, doc, named


def _install_sets(found=()):
    """A fake pynavis.sets inside a FakeNav block (the shared fake has none)."""
    import sys
    import types
    module = types.ModuleType('pynavis.sets')
    module.find = lambda path, doc=None: object() if path in found else None
    sys.modules['pynavis.sets'] = module
    sys.modules['pynavis'].sets = module
    sys.modules['pynavis.doc'].get_filename = lambda d=None: getattr(d, 'FileName', '')


def _remove_sets():
    import sys
    sys.modules.pop('pynavis.sets', None)


def test_gather_facts_through_fake_layer():
    nav, doc, named = _fake_doc()
    profile = _all_on(required_sets=['QA/Hidden'])
    with nav.FakeNav(doc):
        _install_sets()
        try:
            seen = []
            facts = preflight.gather_facts(doc, profile, progress=lambda f, label: seen.append((f, label)))
        finally:
            _remove_sets()
    assert facts['doc_file'] == r'C:\Jobs\Hospital.nwf' and facts['modified'] is True
    assert [m['units'] for m in facts['models']] == ['Meters', 'Millimeters']
    assert [r['name'] for r in facts['stale']] == ['ARCH.rvt', 'MEP.rvt']
    assert facts['sets_found'] == {'QA/Hidden': False}
    assert facts['viewpoints'] == 1
    assert facts['section_on'] is True
    assert facts['clash'] == {'available': True, 'tests': []}
    assert facts['new_clashes']['available'] and facts['new_clashes']['count'] == 0
    assert facts['hidden']['found'] and facts['hidden']['example'].endswith('Hidden box')
    assert facts['overrides']['readable'] is False
    assert facts['timeliner'] == {'available': False}
    assert seen[-1] == (1.0, 'Done') and all(0.0 <= f <= 1.0 for f, _l in seen)
    checks = _by_id(preflight.evaluate(facts, profile))
    assert checks['saved']['level'] == 'warn' and checks['units']['level'] == 'warn'
    assert checks['sets']['level'] == 'fail' and checks['clash_run']['level'] == 'warn'
    assert checks['hidden']['level'] == 'warn' and checks['section']['level'] == 'warn'


def test_run_profile_counts_new_clashes():
    nav, doc, named = _fake_doc()
    results = [nav.ClashResult('Clash1', status='New'), nav.ClashResult('Clash2', status='Active'),
               nav.ClashResult('Clash3', status='New')]
    test = nav.ClashTest('Ducts vs Pipes', [nav.ClashResultGroup('G', results[:2]), results[2]])
    test.LastRun = None
    doc.clash_tests = [test]
    profile = _all_on()
    with nav.FakeNav(doc) as fake:
        _install_sets()
        fake.clash.walk_results = lambda t: _flatten(t.Children, nav)
        try:
            checks = _by_id(preflight.run_profile(profile, doc))
        finally:
            _remove_sets()
    assert checks['new_clashes']['detail'].startswith('2 New clashes')
    assert checks['clash_run']['detail'].startswith('1 test never run (Ducts vs Pipes)')


def _flatten(children, nav):
    for child in children:
        if isinstance(child, nav.ClashResultGroup):
            for leaf in _flatten(child.Children, nav):
                yield leaf
        else:
            yield child


def test_report_writes_html_and_xlsx():
    nav, doc, named = _fake_doc()
    folder = _compat.tmpdir('preflight_report' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        from namo import reports
        with nav.FakeNav(doc):
            _install_sets()
            try:
                result = reports.run('preflight', {'formats': ['html', 'xlsx'], 'folder': folder,
                                                   'name_template': '{report} test'}, doc=doc)
                missing = reports.run('preflight', {'profile': 'Nope', 'folder': folder}, doc=doc)
            finally:
                _remove_sets()
    finally:
        paths.redirect(previous)
    assert result['level'] == 'warning', result
    assert len(result['paths']) == 2 and all(os.path.isfile(p) for p in result['paths'])
    assert result['profile'] == 'Standard' and result['checks']
    assert missing['level'] == 'error' and 'Nope' in missing['message']
    html = open(result['paths'][0], 'rb').read().decode('utf-8')
    assert 'Preflight' in html and 'Sectioning is on' in html
    if not _compat.IRONPYTHON:
        import openpyxl
        wb = openpyxl.load_workbook(result['paths'][1])
        ws = wb['Preflight']
        values = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
        assert 'Sectioning off' in values


def test_report_describe_and_defaults():
    from namo import rep_preflight
    config = rep_preflight.default_config()
    assert config['profile'] == 'Standard' and config['formats'] == ['html'] and 'folder' in config
    assert rep_preflight.describe(config) == 'Profile Standard, html'
    assert rep_preflight.describe({'formats': 'xlsx'}) == 'Profile Standard, xlsx'


# ---- the profile editor form -------------------------------------------------------------------


def test_profile_form_round_trip():
    from namo import ui
    base = preflight.profile_from_choices('QA', ['saved', 'units', 'hidden'], ['saved'], ['ARCH*'],
                                          ['QA/Hidden'], 3, 2)
    fields = ui.normalize_fields(preflight.profile_form_fields(base, True, 'Delete this profile'))
    keys = [f['key'] for f in fields]
    assert keys == ['name', 'enabled', 'blocking', 'expected', 'sets', 'max_new', 'min_views', 'delete']
    raws = {'name': 'QA 2', 'enabled': ['saved', 'hidden'], 'blocking': ['hidden'],
            'expected': 'ARCH*\nMEP', 'sets': '', 'max_new': '-1', 'min_views': '', 'delete': False}
    values, errors = ui.collect(fields, raws)
    assert errors == {}
    profile = preflight.profile_from_form(values, base)
    assert profile['name'] == 'QA 2'
    assert preflight.enabled_checks(profile) == ['saved', 'hidden']
    assert preflight.blocking_checks(profile) == ['hidden']
    assert profile['expected_models'] == ['ARCH*', 'MEP'] and profile['required_sets'] == []
    assert profile['max_new_clashes'] == -1 and profile['min_viewpoints'] == 0
    values, errors = ui.collect(fields, dict(raws, max_new='-5', name='  '))
    assert 'max_new' in errors and 'name' in errors
    plain = ui.normalize_fields(preflight.profile_form_fields(base))
    assert 'delete' not in [f.get('key') for f in plain]
