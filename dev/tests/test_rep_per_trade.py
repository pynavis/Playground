# -*- coding: utf-8 -*-
"""rep_per_trade: who is responsible for each clash, the per-trade plan, file names, the
index workbook and run() into a folder through the fake Navisworks layer."""
import datetime
import os

import _compat

import test_clashreport as fx
from namo import clashreport as cr
from namo import reports
from namo import rep_per_trade as rep

PRIORITY = ['Structure', 'Architecture', 'Mechanical', 'Plumbing', 'Electrical', 'Fire Protection']


def test_contract():
    assert reports.title_of('per_trade') == rep.TITLE
    config = rep.default_config()
    assert config['priority'] == [] and config['subfolder'] is True and config['html'] is False
    assert rep.describe(config) == ('Every test with results, open clashes (New, Active, Reviewed), '
                                    'one workbook per trade')


def test_assigner_settings_reads_clash_assigner():
    stored = {'priority': 'Structure > Mechanical', 'names': 'Mechanical = Jane Smith\nnonsense'}
    calls = []

    def loader(key, defaults):
        calls.append(key)
        return stored
    priority, people = rep.assigner_settings(loader=loader)
    assert calls == [rep.ASSIGNER_KEY]
    assert priority == ['Structure', 'Mechanical'] and people == {'mechanical': 'Jane Smith'}
    priority, people = rep.assigner_settings(['Plumbing'], 'Plumbing = Raj', loader=loader)
    assert priority == ['Plumbing'] and people == {'plumbing': 'Raj'}
    priority, _people = rep.assigner_settings(loader=lambda key, defaults: 1 / 0)
    assert priority[0] == 'Structure'


def test_assigner_key_matches_the_bundle_path():
    """The key pynavis.script.get_config derives for Clash Assigner's folder."""
    import re
    bundle = 'Coordinate.tab/Clash_Results.panel/Clash_Assigner.pushbutton'
    assert rep.ASSIGNER_KEY == 'cfg_' + re.sub(r'[^a-z0-9]+', '_', bundle.lower()).strip('_')
    here = os.path.dirname(os.path.abspath(__file__))
    folder = os.path.join(os.path.dirname(os.path.dirname(here)), 'Playground.extension', *bundle.split('/'))
    assert os.path.isdir(folder)


def test_responsible_rules():
    lookup = rep.people_lookup({'mechanical': 'Jane Smith'}, ['Mechanical', 'Structure'])
    assert lookup == {'jane smith': 'Mechanical'}
    known = ['Mechanical', 'Structure', 'Plumbing']

    def who(**row):
        return rep.responsible(row, PRIORITY, lookup, known)
    assert who(assigned='jane smith') == (['Mechanical'], rep.WHY_ASSIGNED)
    assert who(assigned='plumbing ') == (['Plumbing'], rep.WHY_ASSIGNED)
    assert who(assigned='Contractor X') == (['Contractor X'], rep.WHY_ASSIGNED)
    assert who(trade_a='Structure', trade_b='Mechanical') == (['Mechanical'], rep.WHY_MOVES)
    assert who(trade_a='Electrical', trade_b='Plumbing') == (['Electrical'], rep.WHY_MOVES)
    assert who(trade_a='Mechanical', trade_b='') == (['Mechanical', 'Unassigned'], rep.WHY_BOTH)
    assert who(trade_a='', trade_b=None) == (['Unassigned'], rep.WHY_BOTH)


def test_plan_routes_every_listed_clash():
    data = fx.fake_data()
    plan = rep.plan(data, cr.OPEN, PRIORITY, {'mechanical': 'Jane Smith'}, data['trades_order'])
    trades = [e['trade'] for e in plan]
    assert 'Raj Patel' in trades and 'Jane Smith' not in trades and 'Mechanical' in trades
    listed = [e['listed'] for e in plan]
    assert listed == sorted(listed, reverse=True)
    rows = cr.all_rows(data)
    open_rows = [r for r in rows if r['open']]
    assert sum(listed) == sum(len(r['responsible']) for r in open_rows)
    for entry in plan:
        assert entry['listed'] == sum(entry['why'].values())
    narrowed = rep.trade_data(data, 'Raj Patel')
    assert narrowed['trade'] == 'Raj Patel'
    assert all(r['assigned'] == 'Raj Patel' for t in narrowed['tests'] for r in t['rows'])
    assert all(t['rows'] for t in narrowed['tests'])


def test_file_names():
    now = datetime.datetime(2026, 10, 5, 9, 0)
    assert rep.file_name(rep.NAME_TEMPLATE, 'Fire Protection', 'xlsx', 'Hospital', now) == \
        'Hospital - Clashes - Fire Protection - 2026-10-05.xlsx'
    assert rep.file_name(rep.FOLDER_TEMPLATE, '', '', '', now) == 'Untitled - Clash reports by trade - 2026-10-05'
    assert rep.file_name('{trade}', 'A/B: C', 'xlsx', 'H', now) == 'A_B_ C.xlsx'


def test_index_workbook_links():
    _compat.require_cpython('openpyxl reads the workbook under CPython')
    import openpyxl
    data = fx.fake_data()
    plan = rep.plan(data, cr.OPEN, PRIORITY, {}, data['trades_order'])
    files = dict((e['trade'], {'xlsx': '%s.xlsx' % e['trade'], 'html': '%s.html' % e['trade']}) for e in plan)
    folder = _compat.tmpdir('rep-per-trade-index')
    path = rep.build_index(data, plan, files, PRIORITY, cr.OPEN).save(os.path.join(folder, 'index.xlsx'))
    ws = openpyxl.load_workbook(path)['Trades']
    assert ws['A3'].value == 'Trade' and ws['A4'].value == plan[0]['trade']
    assert ws['I4'].hyperlink.target == '%s.xlsx' % plan[0]['trade']
    assert ws.auto_filter.ref == 'A3:J%d' % (3 + len(plan))
    notes = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any(n and 'Structure > Architecture' in n for n in notes)


def test_run_writes_a_folder():
    doc, _a, _b = fx.fake_clash_doc()
    with fx.Redirected('rep-per-trade-run') as folder:
        with fx.nav.FakeNav(doc):
            result = reports.run('per_trade', {'folder': folder, 'html': True, 'images': False,
                                               'priority': PRIORITY, 'names': 'Architecture = Jane Smith'}, doc=doc)
            nothing = reports.run('per_trade', {'folder': folder, 'statuses': ['Approved'], 'tests': ['Ducts vs Walls'],
                                                'priority': PRIORITY, 'names': 'x = y'}, doc=doc)
    assert result['level'] == 'success', result
    assert result['message'] == 'Wrote clash reports for 2 trades'
    out = result['folder']
    assert os.path.basename(out).startswith('Untitled - Clash reports by trade - ')
    names = sorted(os.listdir(out))
    assert len(names) == 5 and any(n.startswith('Untitled - Clashes by trade') for n in names)
    assert any('Architecture' in n and n.endswith('.html') for n in names)
    assert result['paths'][0].endswith('.xlsx') and 'Clashes by trade' in result['paths'][0]
    assert nothing['level'] == 'info' and nothing['message'] == 'No clashes to send to trades'
