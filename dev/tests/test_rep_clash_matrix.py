# -*- coding: utf-8 -*-
"""rep_clash_matrix: matrices, output-window heat tables and the workbook's graded fills,
plus run() through the fake Navisworks layer."""
import os

import _compat

import test_clashreport as fx
from namo import clashreport as cr
from namo import reports
from namo import rep_clash_matrix as rep


def test_contract():
    assert reports.title_of('clash_matrix') == rep.TITLE
    config = rep.default_config()
    assert config['axis'] == 'trade' and config['xlsx'] is True
    assert rep.describe(config) == 'Every test with results, trade against trade'
    assert rep.describe({'axis': 'model', 'tests': ['A']}) == 'A, model against model'


def test_build_shares_axes_and_counts():
    data = fx.fake_data()
    built = rep.build(data, 'trade')
    open_m, new_m = built['open'], built['new']
    assert open_m['labels'] == new_m['labels']
    assert open_m['labels'][:3] == ['Architecture', 'Structure', 'Mechanical']      # configured order
    assert open_m['total'] == 420 and new_m['total'] == cr.kpis(data)['by_status']['New']
    rows = cr.all_rows(data)
    mech = open_m['labels'].index('Mechanical')
    struct = open_m['labels'].index('Structure')
    expected = sum(1 for r in rows if r['open'] and set([r['trade_a'], r['trade_b']]) == set(['Mechanical',
                                                                                              'Structure']))
    assert open_m['grid'][min(mech, struct)][max(mech, struct)] == expected
    assert built['pairs'][0][:2] == ('Mechanical', 'Structure')
    models = rep.build(data, 'model')
    assert all(l.endswith('.nwc') for l in models['open']['labels'])
    assert rep.headline(built)[0] == '420 clashes open across 6 trades'
    assert rep.headline(built)[1].startswith('Most: Mechanical with Structure (162 open)')


def test_output_html_heat_cells():
    data = fx.fake_data()
    built = rep.build(data, 'trade')
    html = rep.output_html(data, built)
    assert 'Clash count matrix: trade against trade' in html
    assert html.count('<table class="pynavis"') == 3
    assert 'color-mix(in srgb, var(--chart-1) 55%, var(--bg))' in html          # the peak cell
    assert 'background:var(--surface)' in html                                     # below the diagonal
    assert 'Top pairs' in html and '<script' not in html
    empty = rep.matrix_html(cr.matrix([], 'trade'), 'Open clashes')
    assert 'No clashes.' in empty


def test_workbook_fills():
    _compat.require_cpython('openpyxl reads the workbook under CPython')
    import openpyxl
    data = fx.fake_data()
    built = rep.build(data, 'trade')
    folder = _compat.tmpdir('rep-clash-matrix')
    path = rep.build_workbook(data, built).save(os.path.join(folder, 'matrix.xlsx'))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Open', 'New', 'Pairs', 'Tests']
    ws = wb['Open']
    labels = built['open']['labels']
    assert [ws.cell(row=3, column=c + 2).value for c in range(len(labels))] == labels
    assert ws.freeze_panes == 'B4'
    fills = set()
    peak_seen = False
    for r in range(len(labels)):
        for c in range(len(labels)):
            cell = ws.cell(row=4 + r, column=2 + c)
            value = built['open']['grid'][r][c]
            if value is None:
                assert cell.fill.fgColor.rgb.endswith('F2F2F2')
            elif value:
                assert cell.value == value
                fills.add(cell.fill.fgColor.rgb)
                peak_seen = peak_seen or value == rep.peak_of(built['open'])
    assert peak_seen and len(fills) >= 3
    assert ws.cell(row=4, column=len(labels) + 2).value == built['open']['involving'][0]
    pairs = wb['Pairs']
    assert pairs.auto_filter.ref.startswith('A3:D')


def test_run_through_fake_navisworks():
    doc, _a, _b = fx.fake_clash_doc()
    with fx.Redirected('rep-clash-matrix-run') as folder:
        with fx.nav.FakeNav(doc):
            result = reports.run('clash_matrix', {'folder': folder, 'axis': 'model'}, doc=doc, interactive=False)
    assert result['level'] == 'success', result
    assert result['message'] == '3 clashes open across 2 models'
    assert result['paths'] and os.path.isfile(result['paths'][0])
