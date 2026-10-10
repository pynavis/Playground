# -*- coding: utf-8 -*-
"""namo.issue_meeting pure half: choosing the agenda, the meeting log, minutes rows, and the
minutes files (HTML + xlsx, checked with openpyxl under CPython)."""
import datetime
import io
import os

import _compat

from namo import issue_meeting, issues

NOW = datetime.datetime(2026, 10, 5, 14, 30, 0)
TODAY = datetime.date(2026, 10, 5)


def _issue(id_, **fields):
    record = issues.new_issue(dict({'title': 'Issue %s' % id_}, **fields), 'amol', NOW)
    record['id'] = id_
    record['number'] = issues.split_id(id_)[1]
    return record


def _rows():
    return [_issue('A-1', priority='Low', trade='Mechanical', due='2026-10-20'),
            _issue('A-2', priority='Critical', trade='Structure', due='2026-10-01'),
            _issue('A-3', status='Closed', trade='Mechanical'),
            _issue('A-4', priority='High', status='In progress', due='2026-10-07')]


def test_filter_fields_and_choose():
    rows = _rows()
    fields = issue_meeting.filter_fields(rows, {'trades': ['Mechanical', 'Gone'], 'due': 'soon'})
    by_key = dict((f['key'], f) for f in fields)
    assert by_key['trades']['default'] == ['Mechanical']
    assert by_key['due']['default'] == 'soon'
    assert by_key['statuses']['default'] == ['Open', 'In progress']
    ids = [i['id'] for i in issue_meeting.choose(rows, {}, TODAY)]
    assert ids == ['A-2', 'A-4', 'A-1']
    ids = [i['id'] for i in issue_meeting.choose(rows, {'order': 'due'}, TODAY)]
    assert ids == ['A-2', 'A-4', 'A-1']
    ids = [i['id'] for i in issue_meeting.choose(rows, {'due': 'overdue'}, TODAY)]
    assert ids == ['A-2']
    ids = [i['id'] for i in issue_meeting.choose(rows, {'trades': ['Mechanical'],
                                                        'statuses': ['Open', 'Closed']}, TODAY)]
    assert ids == ['A-1', 'A-3']


def test_log_records_net_changes():
    log = issue_meeting.new_log(['A-1', 'A-2', 'A-4'], 'amol', NOW, 'Tower.nwf')
    assert issue_meeting.step(log, 1) == 1 and issue_meeting.step(log, 5) == 2
    assert issue_meeting.step(log, -9) == 0
    assert not issue_meeting.has_activity(log)
    issue_meeting.record(log, 'A-1', 'visit')
    assert not issue_meeting.has_activity(log)
    issue_meeting.record(log, 'A-2', 'status', 'Open', 'In progress', now=NOW)
    issue_meeting.record(log, 'A-2', 'status', 'In progress', 'Resolved', now=NOW)
    issue_meeting.record(log, 'A-2', 'decision', text='Lower the duct 150 mm', now=NOW)
    issue_meeting.record(log, 'A-4', 'due', '2026-10-07', '2026-10-14', now=NOW)
    issue_meeting.record(log, 'A-4', 'assignee', '', 'Ken', now=NOW)
    issue_meeting.record(log, 'A-1', 'discussed', now=NOW)
    issue_meeting.record(log, 'A-1', 'screenshot', now=NOW)
    c = issue_meeting.counts(log)
    assert c == {'issues': 3, 'discussed': 3, 'decisions': 1, 'status_changes': 1,
                 'reassigned': 1, 'new_due': 1, 'screenshots': 1}
    assert log['entries']['A-2']['status'] == ('Open', 'Resolved')
    by_id = dict((i['id'], i) for i in _rows())
    rows = issue_meeting.minutes_rows(log, by_id)
    assert rows[1] == ['A-2', 'Issue A-2', 'Open', 'Open > Resolved', 'Lower the duct 150 mm', '', '', True]
    assert rows[2][5] == 'none > Ken' and rows[2][6] == '2026-10-07 > 2026-10-14'
    assert issue_meeting.duration_text(dict(log, ended='2026-10-05T15:45:00')) == '1 h 15 min'
    assert issue_meeting.duration_text(dict(log, ended='2026-10-05T14:42:00')) == '12 min'


def test_minutes_files():
    log = issue_meeting.new_log(['A-1', 'A-2'], 'amol', NOW, 'Tower.nwf')
    issue_meeting.record(log, 'A-2', 'decision', text='Move the riser <east>', now=NOW)
    issue_meeting.record(log, 'A-2', 'status', 'Open', 'Resolved', now=NOW)
    by_id = dict((i['id'], i) for i in _rows())
    folder = _compat.tmpdir('g8-minutes')
    html_path, xlsx_path = issue_meeting.write_minutes(log, by_id, folder, NOW)
    assert os.path.basename(html_path) == 'Tower - Meeting minutes - 2026-10-05.html'
    with io.open(html_path, encoding='utf-8') as handle:
        page = handle.read()
    assert 'Move the riser &lt;east&gt;' in page and 'Meeting minutes' in page
    again = issue_meeting.write_minutes(log, by_id, folder, NOW)
    assert again[0].endswith('(2).html')
    assert issue_meeting.minutes_base_name('') == issue_meeting.minutes_base_name('Model')
    if _compat.IRONPYTHON:
        return
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    assert wb.sheetnames == ['Minutes', 'Change log']
    sheet = wb['Minutes']
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=4, max_row=6)]
    assert values[0][:2] == ['Issue', 'Title']
    assert values[2][0] == 'A-2' and values[2][3] == 'Open > Resolved'
    assert wb['Change log'].max_row == 3


def test_minutes_folder():
    folder = _compat.tmpdir('g8-minutes-folder')
    assert issue_meeting.minutes_folder(os.path.join(folder, 'Tower.nwf')) == folder
    assert issue_meeting.minutes_folder('').endswith('Documents')
