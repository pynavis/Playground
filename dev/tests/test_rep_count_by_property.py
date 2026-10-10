# -*- coding: utf-8 -*-
"""rep_count_by_property: contract, pivot rows, chart data, and a non-interactive run
writing a workbook through the fake layer."""
import os

import _compat
import test_propindex as pf

from namo import propindex, reports, rep_count_by_property as rep


def test_contract():
    config = rep.default_config()
    assert 'folder' in config and 'name_template' in config
    assert rep.TITLE == reports.title_of('count_by_property')
    assert rep.describe(config) == 'Element > Category, whole model'
    two = dict(config, category2='Element', property2='Level', scope='selection')
    assert rep.specs_of(two) == [('Element', 'Category'), ('Element', 'Level')]
    assert rep.describe(two) == 'Element > Category by Element > Level, current selection'
    assert rep.headers_of(two) == ['Category', 'Level', 'Count', 'Share']
    same = dict(config, category2='Item', property2='Category')
    assert rep.headers_of(same)[:2] == ['Element > Category', 'Item > Category']


def test_rows_and_chart():
    pv = propindex.pivot([('Walls',), ('Ducts',), ('Walls',)])
    assert rep.table_rows(pv)[0][:2] == ['Walls', 2]
    assert _compat.approx(rep.table_rows(pv)[0][2], 2.0 / 3)
    assert rep.chart_data(pv) == (['Walls', 'Ducts'], [2, 1])
    two = propindex.pivot([('Walls', 'L1'), ('Walls', 'L2'), ('Ducts', 'L1')])
    assert rep.chart_data(two) == (['Walls', 'Ducts'], [2, 1])


def test_run_writes_workbook_without_dialogs():
    folder = _compat.tmpdir('g2_rep_count')
    doc, n = pf.tree()
    with pf.Fake(doc):
        path = os.path.join(folder, 'count.xlsx')
        result = rep.run(dict(rep.default_config(), path=path, category2='Element', property2='Mark'),
                         doc=doc, interactive=False)
        assert result['level'] == 'success', result
        assert result['message'] == 'Counted 4 pairs over 4 objects'
        assert result['paths'] == [path] and os.path.exists(path)
        assert 'groups' not in result
        empty = rep.run(dict(rep.default_config(), scope='selection', path=path), doc=doc, interactive=False)
        assert empty['level'] == 'info' and empty['message'] == 'Nothing selected'
        missing = rep.run(dict(rep.default_config(), property=''), doc=doc, interactive=False)
        assert missing['level'] == 'warning'
    try:
        import openpyxl
    except ImportError:
        return
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Counts', 'Totals']
    ws = wb['Counts']
    assert [c.value for c in ws[3]] == ['Category', 'Mark', 'Count', 'Share']
    assert ws.cell(4, 1).value == 'Walls' and ws.cell(4, 3).value == 1


def test_run_through_registry_cancel_is_quiet():
    doc, n = pf.tree()
    with pf.Fake(doc):
        def progress(fraction, label):
            raise reports.Cancelled()
        result = reports.run('count_by_property', {}, doc=doc, interactive=False, progress=progress)
        assert result.get('cancelled') is True


def test_interactive_run_behind_progress_leaves_its_page_for_later():
    """pf.Fake has no pynavis.output: printing during the run would raise. The page
    comes back in result['output'] and prints only when shown."""
    doc, n = pf.tree()
    with pf.Fake(doc):
        result = rep.run(dict(rep.default_config(), category2='Element', property2='Mark'),
                         doc=doc, interactive=True, progress=lambda fraction, label: True)
    assert result['level'] == 'success' and result['groups']
    page = result['output']
    assert isinstance(page, reports.Page)
    kinds = [c[0] for c in page.calls]
    assert kinds[0] == 'html' and 'table' in kinds
    table = [c for c in page.calls if c[0] == 'table'][0]
    assert table[2] == ['Category', 'Mark', 'Count', 'Share'] and len(table[1]) == 4
