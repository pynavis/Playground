# -*- coding: utf-8 -*-
"""Session Stats (namo.mdl_stats + rep_session_stats + hk_sessionstats), Stale File Alert
(rep_stale_files + hk_stalecheck) and the transform-preset hook, pure parts and API
parts through the fake document."""
import io
import json
import os

import _compat

from namo import mdl_stats


def test_sampler_exact_under_budget():
    s = mdl_stats.PrimitiveSampler(budget=10)
    for i in range(10):
        s.add(lambda: (5, 1))
    assert s.total() == (50, 10) and not s.sampled and s.read_count == 10


def test_sampler_doubles_stride_and_extrapolates():
    s = mdl_stats.PrimitiveSampler(budget=100)
    reads = [0]

    def read():
        reads[0] += 1
        return (12, 2)
    for i in range(10000):
        s.add(read)
    primitives, fragments = s.total()
    assert s.sampled and reads[0] < 1000
    assert abs(primitives - 120000) <= 12 * 64 and abs(fragments - 20000) <= 2 * 64
    assert s.count == 10000 and s.read_count == reads[0]


def test_sampler_counts_failures():
    s = mdl_stats.PrimitiveSampler(budget=5)

    def boom():
        raise RuntimeError('no geometry')
    s.add(boom)
    s.add(lambda: (3, 1))
    assert s.failed == 1 and s.total() == (3, 1)


def test_texts_and_totals():
    assert mdl_stats.size_text(((0, 0, 0), (84.24, 36, 52.1)), 'Meters') == '84.2 x 36.0 x 52.1 m'
    assert mdl_stats.size_text(None, 'Meters') == ''
    assert mdl_stats.load_text(None) == 'not recorded'
    assert mdl_stats.load_text({'load_seconds': 12.345}) == '12.3 s'
    assert mdl_stats.arrived_text({'appended_at': None, 'via': 'opened with the document'}) == \
        'opened with the document'
    assert '(Batch Append)' in mdl_stats.arrived_text({'appended_at': 1000000000.0, 'via': 'Batch Append'})
    rows = [{'items': 10, 'geometry': 4, 'primitives': 100, 'fragments': 5, 'cache_size': 2048,
             'sampled': False},
            {'items': 5, 'geometry': 5, 'primitives': 50, 'fragments': 5, 'cache_size': None,
             'sampled': True, 'read': 2}]
    total = mdl_stats.totals(rows)
    assert total == {'models': 2, 'items': 15, 'geometry': 9, 'primitives': 150, 'fragments': 10,
                     'cache_bytes': 2048, 'sampled': True}
    assert mdl_stats.estimate_note(rows[1]) == 'estimated from 2 of 5 geometry items'
    assert mdl_stats.estimate_note(rows[0]) == ''


# ---- through the fake document ----------------------------------------------------------------

import test_items as nav  # noqa: E402


class _Geometry(object):
    PrimitiveCount = 12
    FragmentCount = 1


class _Item(nav.ModelItem):
    @property
    def Geometry(self):
        return _Geometry()


class _Model(object):
    def __init__(self, path, root, source=None):
        self.FileName = path
        self.SourceFileName = source if source is not None else ''
        self.Units = 'Meters'
        self.Creator = 'Revit'
        self.Guid = 'g-' + os.path.basename(path)
        self.RootItem = root
        self.Transform = nav._Ns(IsIdentity=lambda: True)


def _write(path, stamp):
    with io.open(path, 'w', encoding='utf-8') as handle:
        handle.write(u'data')
    os.utime(path, (stamp, stamp))


def _doc(folder):
    arch = os.path.join(folder, 'ARCH.nwc')
    rvt = os.path.join(folder, 'ARCH.rvt')
    _write(arch, 1000000000)
    _write(rvt, 1000090000)
    leaves = [_Item('part %d' % i, geometry=True, box=((i, 0, 0), (i + 1, 2, 3))) for i in range(6)]
    root = nav.ModelItem('ARCH', leaves)
    other = nav.ModelItem('MEP', [_Item('duct', geometry=True, box=((0, 0, 0), (1, 1, 1)))])
    doc = nav.Document([_Model(arch, root, rvt), _Model(os.path.join(folder, 'MEP.nwc'), other)])
    doc.FileName = os.path.join(folder, 'Project.nwf')
    doc.Title = 'Project'
    return doc


def test_hook_stamps_and_collect():
    from namo import hk_sessionstats, session
    session.clear(hk_sessionstats.HOLDER)
    folder = _compat.tmpdir('mdl_stats_collect' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    doc = _doc(folder)
    with nav.FakeNav(doc):
        hk_sessionstats.handle('doc-opened', {'name': 'doc-opened'})
        doc.Models.models.append(_Model(os.path.join(folder, 'STR.nwc'),
                                        nav.ModelItem('STR', [_Item('beam', geometry=True)])))
        hk_sessionstats.handle('model-appended', {'name': 'model-appended'})
        stamps = hk_sessionstats.stamps_for(doc)
        vias = sorted(s['via'] for s in stamps.values())
        assert vias == ['appended', 'opened with the document', 'opened with the document']
        labels = []
        rows = mdl_stats.collect(doc, budget=3, progress=lambda f, label: labels.append(label))
        assert [r['name'] for r in rows] == ['ARCH.rvt', 'MEP.nwc', 'STR.nwc']
        arch = rows[0]
        assert arch['items'] == 7 and arch['geometry'] == 6
        assert arch['sampled'] and arch['read'] < 6
        assert rows[1]['primitives'] == 12 and not rows[1]['sampled']
        assert arch['size'] == '6.0 x 2.0 x 3.0 m'
        assert arch['cache_size'] == 4 and arch['load'] == 'not recorded'
        assert labels[0] == 'ARCH.rvt (1 of 3)' and labels[-1] == 'Done'
        json.dumps(rows)
    session.clear(hk_sessionstats.HOLDER)


def test_record_batch_load_seconds():
    from namo import hk_sessionstats, modelmgmt, session
    session.clear(hk_sessionstats.HOLDER)
    folder = _compat.tmpdir('mdl_stats_batch' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    doc = _doc(folder)
    with nav.FakeNav(doc):
        key = modelmgmt.models_info(doc)[1]['key']
        hk_sessionstats.record_batch(doc, {key: 3.25})
        rows = mdl_stats.collect(doc)
        assert rows[1]['load'] == '3.2 s' or rows[1]['load'] == '3.3 s'
        assert 'Batch Append' in rows[1]['arrived']
    session.clear(hk_sessionstats.HOLDER)


def test_reports_contract_and_files():
    from namo import paths, reports, rep_session_stats, rep_stale_files
    folder = _compat.tmpdir('mdl_stats_reports' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        for module in (rep_session_stats, rep_stale_files):
            config = module.default_config()
            assert 'folder' in config and 'name_template' in config
            json.dumps(config, sort_keys=True)
            assert module.describe(config)
            assert module.TITLE == reports.title_of(module.__name__.split('rep_', 1)[1])
        doc = _doc(folder)
        with nav.FakeNav(doc):
            out = reports.run('stale_files', {'folder': folder}, doc=doc)
            assert out['level'] == 'warning' and out['message'] == '1 model out of date', out
            assert os.path.isfile(out['paths'][0]) and len(out['rows']) == 2
            statuses = [r['status'] for r in out['rows']]
            assert statuses == ['Out of date', 'File not found']
            out = reports.run('session_stats', {'folder': folder, 'budget': 2}, doc=doc)
            assert out['level'] == 'success' and out['totals']['items'] == 9, out
            assert out['message'] == '2 models, 84 primitives' or out['totals']['sampled']
            assert os.path.isfile(out['paths'][0])
            empty = nav.Document([])
            out = reports.run('session_stats', {'folder': folder}, doc=empty)
            assert out['level'] == 'info' and out['paths'] == []
    finally:
        paths.redirect(previous)


def test_stale_workbook_round_trip():
    _compat.require_cpython('openpyxl')
    import openpyxl
    from namo import paths, reports
    folder = _compat.tmpdir('mdl_stats_xlsx' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        doc = _doc(folder)
        with nav.FakeNav(doc):
            out = reports.run('stale_files', {'folder': folder}, doc=doc)
        book = openpyxl.load_workbook(out['paths'][0])
        sheet = book['Files']
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        header = [v for v in values if v and v[0] == 'Model'][0]
        assert header[:3] == ['Model', 'Loaded file', 'File time']
        assert any(v and v[6] == 'Out of date' for v in values)
    finally:
        paths.redirect(previous)


def test_stalecheck_hook_quiet_by_default():
    from namo import hk_stalecheck
    doc = _doc(_compat.tmpdir('mdl_stats_hook' + ('_ipy' if _compat.IRONPYTHON else '_cpy')))
    with nav.FakeNav(doc):
        assert hk_stalecheck.load_config()['check_on_open'] is False
        hk_stalecheck.handle('doc-opened', {'name': 'doc-opened'})
        hk_stalecheck.handle('model-appended', {'name': 'model-appended'})


def test_transformpresets_hook_without_presets():
    from namo import hk_transformpresets, paths
    folder = _compat.tmpdir('mdl_stats_presets' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        doc = _doc(folder)
        with nav.FakeNav(doc):
            hk_transformpresets.handle('model-appended', {'name': 'model-appended'})
            hk_transformpresets.handle('doc-opened', {'name': 'doc-opened'})
    finally:
        paths.redirect(previous)
