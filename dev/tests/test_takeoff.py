# -*- coding: utf-8 -*-
"""namo.takeoff: property detection, SI conversion, grouping, zones, pivots, workbooks."""
import math
import os

import _compat
from namo import takeoff as tk


def test_choose_pairs_prefers_element_and_net():
    sampled = [('Element', 'Volume', 40), ('Qto_WallBaseQuantities', 'GrossVolume', 12),
               ('Qto_WallBaseQuantities', 'NetVolume', 12), ('Element', 'Length', 30),
               ('Revit Type', 'Length', 2), ('Custom', 'Length', 50), ('Element', 'Area', 40),
               ('BaseQuantities', 'NetSideArea', 5), ('Element', 'Mark', 40), ('Data', 'Weight', 3)]
    pairs = tk.choose_pairs(sampled)
    assert pairs['volume'] == [('Element', 'Volume'), ('Qto_WallBaseQuantities', 'NetVolume'),
                               ('Qto_WallBaseQuantities', 'GrossVolume')]
    assert pairs['length'] == [('Element', 'Length'), ('Revit Type', 'Length'), ('Custom', 'Length')]
    assert pairs['area'] == [('Element', 'Area'), ('BaseQuantities', 'NetSideArea')]
    assert pairs['weight'] == [('Data', 'Weight')]
    assert tk.choose_pairs([], ['volume']) == {'volume': []}


def test_to_si_numbers_and_display_strings():
    assert tk.to_si('length', 10.0, 'Feet') == 10 * 0.3048
    assert _compat.approx(tk.to_si('area', 10.0, 'Feet'), 10 * 0.3048 ** 2)
    assert _compat.approx(tk.to_si('volume', 2.0, 'Millimeters'), 2e-9)
    assert tk.to_si('length', '12\' - 6"', 'Meters') == _approx_m(12.5 * 0.3048)
    assert _compat.approx(tk.to_si('area', '250 SF', 'Meters'), 250 * 0.09290304)
    assert _compat.approx(tk.to_si('volume', '3.5 m³', 'Feet'), 3.5)
    assert tk.to_si('volume', '3.5 m2', 'Feet') is None            # wrong kind
    assert tk.to_si('weight', 100, 'Feet') == 45.359237
    assert tk.to_si('weight', '1.2 t', 'Feet') == 1200.0
    assert tk.to_si('weight', 7, 'Meters') == 7.0
    assert tk.to_si('length', '2.5', 'Meters') == 2.5
    assert tk.to_si('length', 'n/a', 'Meters') is None
    assert tk.to_si('length', None, 'Meters') is None and tk.to_si('length', True, 'Meters') is None
    assert tk.to_si('length', float('nan'), 'Meters') is None


def _approx_m(value):
    return float('%.12g' % value)


def _record(groups, zone=None, **measures):
    sources = dict((m, 'property') for m in measures if measures[m] is not None)
    if 'src' in measures:
        sources = measures.pop('src')
    return {'groups': list(groups), 'zone': zone, 'measures': measures, 'sources': sources,
            'name': '/'.join(groups)}


def test_aggregate_one_and_two_levels():
    records = [_record(['Walls', 'L1'], volume=2.0, length=4.0),
               _record(['Walls', 'L2'], volume=3.0, length=None),
               _record(['Ducts', 'L1'], volume=None, length=10.0),
               _record(['Walls', 'L1'], volume=1.0, length=1.0, src={'volume': 'geometric', 'length': 'property'})]
    rows = tk.aggregate(records, 1, ['count', 'volume', 'length'])
    assert [(r['kind'], r['keys'], r['count']) for r in rows] == [
        ('group', ['Ducts'], 1), ('group', ['Walls'], 3), ('total', ['Total'], 4)]
    walls = rows[1]
    assert walls['values'] == {'volume': 6.0, 'length': 5.0}
    assert walls['missing'] == {'volume': 0, 'length': 1}
    assert walls['sources']['volume'] == {'property': 2, 'geometric': 1}
    assert tk.source_note(walls, 'volume') == '1 geometric'
    assert tk.source_note(walls, 'length') == '1 without'
    assert rows[0]['values']['volume'] is None and rows[0]['missing']['volume'] == 1
    two = tk.aggregate(records, 2, ['count', 'volume'])
    assert [(r['kind'], r['keys'], r['count']) for r in two] == [
        ('group', ['Ducts', 'L1'], 1), ('subtotal', ['Ducts', 'Subtotal'], 1),
        ('group', ['Walls', 'L1'], 2), ('group', ['Walls', 'L2'], 1),
        ('subtotal', ['Walls', 'Subtotal'], 3), ('total', ['Total', ''], 4)]
    assert two[4]['values']['volume'] == 6.0


def test_natural_group_order_and_texts():
    records = [_record([name]) for name in ('Level 10', 'Level 2', '(no value)', 'level 1')]
    rows = tk.aggregate(records, 1, ['count'])
    assert [r['keys'][0] for r in rows[:-1]] == ['(no value)', 'level 1', 'Level 2', 'Level 10']
    assert tk.group_text(None) == '(no value)' and tk.group_text('  ') == '(no value)'
    assert tk.group_text(3.0) == '3' and tk.group_text(True) == 'Yes'


def test_display_units_and_rows():
    unit_set = tk.display_units('auto', 'Feet')
    assert unit_set == {'length': 'ft', 'area': 'ft2', 'volume': 'ft3', 'surface': 'ft2', 'weight': 'lb'}
    assert tk.display_units('mm', 'Feet')['volume'] == 'm3'
    assert tk.header('volume', unit_set) == 'Volume (ft³)' and tk.header('count', unit_set) == 'Count'
    rows = tk.aggregate([_record(['Walls'], volume=1.0)], 1, ['count', 'volume'])
    table = tk.table_rows(rows, ['count', 'volume'], tk.display_units('m', 'Meters'), ['Category'])
    assert table[0] == {'cells': ['Walls', 1, 1.0, ''], 'value': 0}
    assert tk.table_headers(['count', 'volume'], unit_set, ['Category']) == [
        'Category', 'Count', 'Volume (ft³)', 'Notes']
    assert _compat.approx(tk.from_si('volume', 1.0, unit_set), 35.3146667, 1e-6)


def test_level_bands():
    bands = tk.level_bands([('Level 2', 4.0), ('Level 1', 0.0), ('Roof', 8.0)])
    assert [b['name'] for b in bands] == ['Below Level 1', 'Level 1', 'Level 2', 'Roof']
    assert tk.assign_band(-1, bands) == 'Below Level 1'
    assert tk.assign_band(0.0, bands) == 'Level 1' and tk.assign_band(3.99, bands) == 'Level 1'
    assert tk.assign_band(4.0, bands) == 'Level 2' and tk.assign_band(100, bands) == 'Roof'
    assert tk.level_bands([]) == []


def test_assign_box_centre_then_overlap():
    zones = [{'name': 'A', 'box': ((0, 0, 0), (10, 10, 3))},
             {'name': 'B', 'box': ((10, 0, 0), (20, 10, 3))},
             {'name': 'Big', 'box': ((0, 0, 0), (20, 10, 3))}]
    assert tk.assign_box(((1, 1, 0), (3, 3, 1)), zones[:2]) == 'A'
    assert tk.assign_box(((8, 1, 0), (16, 3, 1)), zones[:2]) == 'B'      # centre x=12
    assert tk.assign_box(((9, 1, 0), (11, 3, 1)), zones[:2]) == 'A'      # on the shared wall: tie, first
    assert tk.assign_box(((1, 1, 0), (3, 3, 1)), zones) == 'A'           # both hold it, A overlaps... equally
    assert tk.assign_box(((30, 1, 0), (31, 3, 1)), zones) == tk.NO_ZONE
    assert tk.assign_box(((-2, 1, 0), (1, 3, 1)), zones[:2]) == 'A'      # centre outside, overlaps A
    assert tk.unique_names(['Office', 'Office', '', 'Office']) == ['Office', 'Office (2)', 'Zone', 'Office (3)']


def test_tiles():
    assert tk.assign_tile((5, 5, 0), (0, 0, 0), 10) == 'X01 Y01'
    assert tk.assign_tile((25, 5, 0), (0, 0, 0), 10, (120, 3)) == 'X003 Y01'
    assert tk.assign_tile((-5, -5, 0), (0, 0, 0), 10) == 'X01 Y01'
    _compat.raises(ValueError, tk.assign_tile, (0, 0, 0), (0, 0, 0), 0)


def _grid(angle_deg=0.0):
    a = math.radians(angle_deg)
    ux, uy = math.cos(a), math.sin(a)
    vx, vy = -math.sin(a), math.cos(a)
    hits = []
    letters = ['A', 'B', 'C']
    numbers = ['1', '2', '3']
    for i, letter in enumerate(letters):         # letter lines run along v, spaced along u
        for j, number in enumerate(numbers):     # number lines run along u, spaced along v
            x = i * 8.0 * ux + j * 6.0 * vx
            y = i * 8.0 * uy + j * 6.0 * vy
            hits.append({'line1': letter, 'line2': number, 'dir1': (vx, vy), 'dir2': (ux, uy),
                         'pos': (x, y)})
    return hits


def test_grid_bays_straight_and_rotated():
    for angle in (0.0, 30.0):
        families = tk.grid_families(_grid(angle))
        assert [[n for _o, n in f['lines']] for f in families] == [['A', 'B', 'C'], ['1', '2', '3']]
        a = math.radians(angle)

        def at(u, v):
            return (u * math.cos(a) - v * math.sin(a), u * math.sin(a) + v * math.cos(a), 0.0)
        assert tk.grid_bay(at(4, 3), families) == 'A-B / 1-2'
        assert tk.grid_bay(at(12, 9), families) == 'B-C / 2-3'
        assert tk.grid_bay(at(-1, 13), families) == '< A / > 3'
    assert tk.grid_families([]) == []
    assert tk.grid_bay((0, 0, 0), []) == tk.NO_ZONE


def test_pivot_and_other_column():
    records = [_record(['Walls'], 'L1', volume=1.0), _record(['Walls'], 'L2', volume=2.0),
               _record(['Ducts'], 'L1', volume=None), _record(['Pipes'], None, volume=4.0)]
    zones, groups, matrix, zone_totals, group_totals = tk.pivot(records, 'count')
    assert zones == ['(no zone)', 'L1', 'L2'] and groups == ['Ducts', 'Pipes', 'Walls']
    assert matrix[('L1', 'Walls')] == 1.0 and zone_totals['L1'] == 2.0 and group_totals['Walls'] == 2.0
    zones, groups, matrix, zone_totals, group_totals = tk.pivot(records, 'volume', top=1)
    assert groups == ['Pipes', 'Other'] and matrix[('L2', 'Other')] == 2.0


def _sheets(path):
    import zipfile
    with zipfile.ZipFile(path) as package:
        text = package.read('xl/workbook.xml').decode('utf-8')
    return [part.split('"')[0] for part in text.split('<sheet name="')[1:]]


def test_quick_and_zone_workbooks_round_trip():
    folder = _compat.tmpdir('takeoff_xlsx')
    records = [_record(['Walls', 'L1'], 'Zone A', volume=2.0, length=4.0),
               _record(['Walls', 'L2'], 'Zone B', volume=3.0, length=1.0),
               _record(['Ducts', 'L1'], 'Zone A', volume=None, length=10.0,
                       src={'length': 'bbox'})]
    unit_set = tk.display_units('m', 'Meters')
    measures = ['count', 'volume', 'length']
    rows = tk.aggregate(records, 2, measures)
    path = tk.write_quick_xlsx(os.path.join(folder, 'quick.xlsx'), 'Quick takeoff', 'Model.nwf',
                               rows, measures, unit_set, ['Category', 'Level'],
                               [('Scope', 'Whole model')], {'volume': [('Element', 'Volume')]}, records)
    assert _sheets(path) == ['Takeoff', 'Objects', 'Notes']
    zpath, written, zones = tk.write_zone_xlsx(os.path.join(folder, 'zone.xlsx'), 'Zone takeoff',
                                               'Model.nwf', records, measures, unit_set,
                                               ['Category'], [('Zones', 'Level bands')], 'volume')
    assert _sheets(zpath) == ['Pivot', 'Zones', 'Zone A', 'Zone B', 'Notes']
    assert written == 2 and zones == ['Zone A', 'Zone B']
    if _compat.IRONPYTHON:
        return
    import openpyxl
    book = openpyxl.load_workbook(path)
    sheet = book['Takeoff']
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=3)]
    assert values[0][:5] == ['Category', 'Level', 'Count', 'Volume (m³)', 'Length (m)']
    assert values[-1][:5] == ['Total', None, 3, 5.0, 15.0]
    assert values[2][:3] == ['Ducts', 'Subtotal', 1]
    assert 'Length: 1 from bounding box' in values[1][5]
    notes = [[c.value for c in row] for row in book['Notes'].iter_rows(min_row=2)]
    assert ['Volume read from', 'Element > Volume'] in notes
    book = openpyxl.load_workbook(zpath)
    pivot = [[c.value for c in row] for row in book['Pivot'].iter_rows(min_row=3)]
    assert pivot[0] == ['Zone', 'Walls', 'Ducts', 'Total'] or pivot[0] == ['Zone', 'Ducts', 'Walls', 'Total']
    assert pivot[-1][0] == 'Total' and pivot[-1][-1] == 5.0


# ---- solids from triangles (regressions from the refit pass) --------------------------------


def _quad(a, b, c, d):
    return [(a, b, c), (a, c, d)]


def _box(x0, y0, z0, x1, y1, z1, skip=(), flip=()):
    """An outward-wound box as 12 triangles; skip/flip name faces to leave out or reverse."""
    p = lambda x, y, z: (float(x), float(y), float(z))  # noqa: E731
    faces = {
        'bottom': _quad(p(x0, y0, z0), p(x0, y1, z0), p(x1, y1, z0), p(x1, y0, z0)),
        'top': _quad(p(x0, y0, z1), p(x1, y0, z1), p(x1, y1, z1), p(x0, y1, z1)),
        'front': _quad(p(x0, y0, z0), p(x1, y0, z0), p(x1, y0, z1), p(x0, y0, z1)),
        'back': _quad(p(x0, y1, z0), p(x0, y1, z1), p(x1, y1, z1), p(x1, y1, z0)),
        'left': _quad(p(x0, y0, z0), p(x0, y0, z1), p(x0, y1, z1), p(x0, y1, z0)),
        'right': _quad(p(x1, y0, z0), p(x1, y1, z0), p(x1, y1, z1), p(x1, y0, z1)),
    }
    out = []
    for name in sorted(faces):
        if name in skip:
            continue
        tris = faces[name]
        if name in flip:
            tris = [(a, c, b) for a, b, c in tris]
        out.extend(tris)
    return out


def test_solid_measure_closed_box():
    solid = tk.solid_measure(_box(0, 0, 0, 2, 3, 4))
    assert _compat.approx(solid['volume'], 24.0)
    assert all(_compat.approx(v, w) for v, w in zip(solid['centroid'], (1.0, 1.5, 2.0)))
    inward = [(a, c, b) for a, b, c in _box(0, 0, 0, 2, 3, 4)]
    assert _compat.approx(tk.solid_measure(inward)['volume'], 24.0)
    assert tk.solid_measure([]) is None


def test_solid_measure_refuses_open_and_miswound_meshes():
    assert tk.solid_measure(_box(0, 0, 0, 2, 3, 4, skip=('top',))) is None
    # An open tube: its area vectors cancel, so only the per-axis volumes give it away.
    assert tk.solid_measure(_box(0, 0, 0, 2, 3, 4, skip=('top', 'bottom'))) is None
    assert tk.solid_measure(_box(0, 0, 0, 2, 3, 4, flip=('top',))) is None
    sheet = _quad((0.0, 0.0, 0.0), (1.0, 0.0, 0.0), (1.0, 1.0, 0.0), (0.0, 1.0, 0.0))
    assert tk.solid_measure(sheet) is None
    assert tk.solid_measure(sheet + [(a, c, b) for a, b, c in sheet]) is None    # two-sided sheet


def test_solid_measure_does_not_need_welded_faces():
    # Regression: a watertight solid whose faces do not share vertices (how Navisworks
    # tessellates neighbouring faces) failed the welded-edge test and was weighed as a
    # bounding box times a fill factor. Here the top is split into four around a middle
    # vertex the side faces never use (T-junctions on every top edge).
    from namo import mesh
    box = _box(0, 0, 0, 2, 2, 1, skip=('top',))
    m = (1.0, 1.0, 1.0)
    mids = [(1.0, 0.0, 1.0), (2.0, 1.0, 1.0), (1.0, 2.0, 1.0), (0.0, 1.0, 1.0)]
    corners = [(0.0, 0.0, 1.0), (2.0, 0.0, 1.0), (2.0, 2.0, 1.0), (0.0, 2.0, 1.0)]
    top = []
    for i in range(4):
        top.append((m, corners[i], mids[i]))
        top.append((m, mids[i], corners[(i + 1) % 4]))
    tris = box + top
    assert not mesh.is_closed(tris)
    solid = tk.solid_measure(tris)
    assert solid is not None and _compat.approx(solid['volume'], 4.0)


def test_solid_measure_far_from_the_origin():
    # The owner's model sits about 1.37 million feet out in shared coordinates.
    far = _box(1370000.0, -1370000.0, 300.0, 1370002.0, -1369997.0, 304.0)
    solid = tk.solid_measure(far)
    assert abs(solid['volume'] - 24.0) < 1e-6
    assert abs(solid['centroid'][0] - 1370001.0) < 1e-6 and abs(solid['centroid'][1] + 1369998.5) < 1e-6
