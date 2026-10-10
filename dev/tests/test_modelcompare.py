# -*- coding: utf-8 -*-
"""namo.modelcompare: the match chain, diff reasons, property hashes, snapshots, the
workbook, and the API half through the fake Navisworks layer."""
import json
import os
import sys
import types

import _compat

from namo import modelcompare as mc


def _rec(uid=None, ifc=None, eid=None, guid=None, name='Duct', cls='Duct', at=0.0, size=1.0, h=None):
    box = ((at, 0.0, 0.0), (at + size, 1.0, 1.0))
    return mc.record(uid, ifc, eid, guid, name, cls, box, h)


def test_record_cleans_values():
    rec = mc.record(' ', '', 12, '{00000000-0000-0000-0000-000000000000}', None, None, None, None)
    assert rec['uid'] is None and rec['ifc'] is None and rec['eid'] == '12' and rec['guid'] is None
    assert rec['name'] == '' and rec['box'] is None
    assert mc.record(guid='ABC')['guid'] == 'ABC'


def test_match_chain_prefers_ids_in_order():
    old = [_rec(uid='U1', eid='1'), _rec(ifc='I2'), _rec(eid='3'), _rec(guid='G4'),
           _rec(name='Plain', at=50.0), _rec(uid='GONE', name='Gone', at=99.0)]
    new = [_rec(guid='G4', at=5.0), _rec(name='Plain', at=50.0), _rec(uid='u1', eid='999', at=1.0),
           _rec(eid='3', at=2.0), _rec(ifc='I2'), _rec(uid='NEW', name='New', at=77.0)]
    pairs, removed, added = mc.match(old, new, 0.001)
    assert pairs == [(0, 2, 'uid'), (1, 4, 'ifc'), (2, 3, 'eid'), (3, 0, 'guid'), (4, 1, 'shape')]
    assert removed == [5] and added == [5]


def test_ambiguous_keys_fall_through():
    # two old walls share an Element ID: the id is not used, their boxes decide
    old = [_rec(eid='7', name='Wall', at=0.0), _rec(eid='7', name='Wall', at=10.0)]
    new = [_rec(eid='7', name='Wall', at=10.0), _rec(eid='7', name='Wall', at=0.0)]
    pairs, removed, added = mc.match(old, new, 0.001)
    assert pairs == [(0, 1, 'shape'), (1, 0, 'shape')] and removed == [] and added == []
    # ambiguous on one side only is ambiguous too
    old = [_rec(eid='8', at=0.0), _rec(eid='8', at=0.0)]
    new = [_rec(eid='8', at=0.0)]
    pairs, removed, added = mc.match(old, new, 0.001)
    assert pairs == [] and removed == [0, 1] and added == [0]


def test_shape_key_rounds_to_the_tolerance():
    a = _rec(name='Pipe', at=1.00004)
    b = _rec(name='PIPE', at=1.00002)
    assert mc.shape_key(a, 0.001) == mc.shape_key(b, 0.001)
    assert mc.shape_key(_rec(name='Pipe', at=1.01), 0.001) != mc.shape_key(a, 0.001)
    assert mc.shape_key(mc.record(name='x'), 0.001) is None


def test_diff_reasons():
    old = [_rec(uid='A'), _rec(uid='B'), _rec(uid='C', h='h1'), _rec(uid='D', h='h1'), _rec(uid='E')]
    new = [_rec(uid='A', at=0.5), _rec(uid='B', size=1.2), _rec(uid='C', h='h2'), _rec(uid='D', h='h1'),
           _rec(uid='E', at=0.0005)]
    pairs, removed, added = mc.match(old, new, 0.001)
    changed, unchanged = mc.diff(old, new, pairs, 0.001)
    reasons = dict((old[i]['uid'], [r[0] for r in rs]) for i, _j, _s, rs in changed)
    assert reasons == {'A': ['moved'], 'B': ['moved', 'resized'], 'C': ['properties']}
    assert sorted(old[i]['uid'] for i, _j, _s in unchanged) == ['D', 'E']
    changed, unchanged = mc.diff(old, new, pairs, 0.001, compare_hash=False)
    assert 'C' in [old[i]['uid'] for i, _j, _s in unchanged]
    stats = mc.counts(pairs, removed, added, *mc.diff(old, new, pairs, 0.001))
    assert stats['changed'] == 3 and stats['moved'] == 2 and stats['resized'] == 1
    assert stats['properties'] == 1 and stats['by_scheme'] == {'uid': 5}


def test_reason_text_and_summary():
    text = mc.reason_text([('moved', 0.12), ('resized', 0.015), ('properties', None)],
                          lambda v: '%d mm' % round(v * 1000))
    assert text == 'Moved 120 mm, resized 15 mm, properties changed'
    assert mc.reason_text([], str) == ''
    stats = {'added': 0, 'removed': 0, 'changed': 0}
    assert mc.summary(stats, 1234) == ('info', 'No differences found', '1,234 objects match.')
    assert 'budget' in mc.summary(stats, 1234, sampled=True)[2]
    level, message, detail = mc.summary({'added': 1200, 'removed': 3, 'changed': 45}, 12345)
    assert level == 'success' and message == '1,200 added, 45 changed, 3 removed'
    assert detail == 'Compared 12,345 objects.'


def test_property_lines_and_hash():
    pairs, problems = mc.parse_property_lines('Element > Category\n\nElement>Type; Bad line\n > x\nElement > Type')
    assert pairs == [('Element', 'Category'), ('Element', 'Type')]
    assert len(problems) == 2 and 'Bad line' in problems[0]
    h1 = mc.property_hash([('Element', 'Type', 'A'), ('Element', 'Category', 'Ducts')])
    h2 = mc.property_hash([('Element', 'Category', 'Ducts'), ('Element', 'Type', 'A')])
    assert h1 == h2 and len(h1) == 16
    assert mc.property_hash([('Element', 'Type', 'B'), ('Element', 'Category', 'Ducts')]) != h1
    assert mc.property_hash([('Element', 'Type', None)]) == mc.property_hash([('Element', 'Type', '')])
    assert mc.property_hash([]) is None
    assert mc.property_hash([(u'Elément', u'Typ\u00e9', u'\u00e9')])


def test_snapshot_round_trip_and_scale():
    records = [_rec(uid='U', ifc='I', eid='1', guid='G', name=u'Dücht', at=1.123456789, h='abc'),
               mc.record(name='No box')]
    snap = mc.snapshot_record('Before', 'ARCH.nwc', 'Meters', records, [('Element', 'Type')],
                              '2026-10-05 14:30', sampled=True)
    text = json.dumps(snap, sort_keys=True)
    back = json.loads(text)
    assert mc.check_snapshot(back) is None
    assert back['count'] == 2 and back['properties'] == [['Element', 'Type']] and back['sampled']
    decoded = mc.decode_records(back['objects'])
    assert decoded[0]['uid'] == 'U' and decoded[0]['name'] == u'Dücht' and decoded[0]['hash'] == 'abc'
    assert _compat.approx(decoded[0]['box'][0][0], 1.1235)
    assert decoded[1]['box'] is None
    scaled = mc.decode_records(back['objects'], scale=1000.0)
    assert _compat.approx(scaled[0]['box'][1][1], 1000.0)
    meta = mc.snapshot_meta(dict(back, id='S-001'))
    assert 'objects' not in meta and meta['id'] == 'S-001'
    assert mc.check_snapshot(None) and mc.check_snapshot(dict(back, version=99))
    assert mc.check_snapshot(dict(back, objects=None))


def test_write_xlsx():
    folder = _compat.tmpdir('modelcompare_xlsx' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    stats = {'added': 1, 'changed': 1, 'removed': 0, 'unchanged': 5, 'moved': 1, 'resized': 0,
             'properties': 0, 'by_scheme': {'uid': 6}}
    path = mc.write_xlsx(os.path.join(folder, 'c.xlsx'), stats,
                         [('Added', ['Object', 'Class'], [['Duct 1', 'Duct']]),
                          ('Changed', ['Object', 'What changed'], [['Pipe', 'Moved 5 mm']]),
                          ('Removed', ['Object'], [])], subtitle='A to B')
    assert os.path.getsize(path) > 0
    _compat.require_cpython('openpyxl round trip')
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Summary', 'Added', 'Changed', 'Removed']
    values = [r for r in wb['Summary'].iter_rows(values_only=True)]
    assert ('Added', 1) in values and ('Matched by Revit UniqueId', 6) in values
    assert list(wb['Changed'].iter_rows(values_only=True))[1] == ('Pipe', 'Moved 5 mm')


# ---- API half through the fake layer ----------------------------------------------------------

import test_items as nav  # noqa: E402


class Item(nav.ModelItem):
    def __init__(self, name, children=(), geometry=False, box=None, props=None, guid=None, klass='Duct'):
        nav.ModelItem.__init__(self, name, children, geometry=geometry, box=box, props=props, klass=klass)
        self.InstanceGuid = guid or '00000000-0000-0000-0000-000000000000'


def _object(name, x, uid=None, eid=None, guid=None, type_name='A'):
    leaf = Item(name + ' solid', geometry=True, box=((x, 0, 0), (x + 1, 1, 1)))
    props = {('Element', 'Type'): type_name}
    if uid:
        props[('Element', 'UniqueId')] = uid
    if eid:
        props[('Element ID', 'Value')] = eid
    element = Item(name, [leaf], props=props, guid=guid)
    leaf.object = element
    return element


def _model(name, objects):
    root = Item(name, objects)
    model = nav.Model('C:\\X\\' + name, root)
    root.Model, root.HasModel = model, True
    return model


def _versions():
    old = _model('V1.nwc', [_object('Kept', 0, uid='K'), _object('Moved', 5, uid='M'),
                            _object('Retyped', 9, uid='R'), _object('Dropped', 20, uid='D')])
    new = _model('V2.nwc', [_object('Kept', 0, uid='K'), _object('Moved', 7, uid='M'),
                            _object('Retyped', 9, uid='R', type_name='B'), _object('Added', 30, uid='N')])
    return nav.Document([old, new])


def test_collect_probes_keys_and_hashes():
    doc = _versions()
    with nav.FakeNav(doc):
        data = mc.collect(doc.Models[0].RootItem, prop_pairs=[('Element', 'Type')])
    assert data['pairs']['uid'] == ('Element', 'UniqueId')
    assert data['pairs']['eid'] is None and data['pairs']['ifc'] is None
    names = [r['name'] for r in data['records']]
    assert names == ['Kept', 'Moved', 'Retyped', 'Dropped'] and not data['sampled']
    assert data['records'][0]['uid'] == 'K' and data['records'][0]['guid'] is None
    assert data['records'][0]['hash'] == mc.property_hash([('Element', 'Type', 'A')])
    assert data['records'][1]['box'] == ((5, 0, 0), (6, 1, 1))


def _fake_sets(log):
    module = types.ModuleType('pynavis.sets')
    existing = {'Compare/Changed': 'changed-set'}

    def find(path, doc=None):
        return existing.get(path)

    def update(item, source, doc=None):
        log.append(('update', item, [i.DisplayName for i in source]))

    def create(name, source, folder=None, doc=None):
        log.append(('create', '%s/%s' % (folder, name), [i.DisplayName for i in source]))
    module.find, module.update, module.create = find, update, create
    return module


def test_compare_two_models_paints_and_writes_sets_in_one_transaction():
    doc = _versions()
    log = []
    with nav.FakeNav(doc):
        sys.modules['pynavis.sets'] = _fake_sets(log)
        sys.modules['pynavis'].sets = sys.modules['pynavis.sets']
        try:
            pairs_pp = [('Element', 'Type')]
            old = mc.collect(doc.Models[0].RootItem, prop_pairs=pairs_pp)
            new = mc.collect(doc.Models[1].RootItem, prop_pairs=pairs_pp)
            pairs, removed, added = mc.match(old['records'], new['records'], 0.001)
            changed, unchanged = mc.diff(old['records'], new['records'], pairs, 0.001)
            assert [new['records'][j]['name'] for j in added] == ['Added']
            assert [old['records'][i]['name'] for i in removed] == ['Dropped']
            reasons = dict((new['records'][j]['name'], [r[0] for r in rs]) for _i, j, _s, rs in changed)
            assert reasons == {'Moved': ['moved'], 'Retyped': ['properties']}
            written, skipped, failed = mc.apply_results(
                [new['objects'][j] for j in added], [new['objects'][j] for _i, j, _s, _r in changed],
                [old['objects'][i] for i in removed], [new['objects'][j] for _i, j, _s in unchanged],
                True, True, doc)
        finally:
            sys.modules.pop('pynavis.sets', None)
    # the outer transaction comes first; paint's own one nests inside it (the fake layer
    # records both, the real pynavis transaction skips the inner one: one undo step)
    assert doc.transactions[0] == 'Compare models' and set(doc.transactions) == set(['Compare models'])
    kinds = [c[0] for c in doc.Models.calls]
    assert kinds == ['transparency', 'color', 'color', 'color']
    assert written == ['Added', 'Changed', 'Removed'] and skipped == [] and failed == []
    assert log == [('create', 'Compare/Added', ['Added']), ('update', 'changed-set', ['Moved', 'Retyped']),
                   ('create', 'Compare/Removed', ['Dropped'])]


def test_write_sets_skips_empty_lists_and_reports_failures():
    doc = _versions()
    with nav.FakeNav(doc):
        module = _fake_sets([])

        def broken(name, source, folder=None, doc=None):
            raise ValueError('nope')
        module.create = broken
        sys.modules['pynavis.sets'] = module
        try:
            written, skipped, failed = mc.write_sets({'Added': [], 'Changed': [doc.Models[0].RootItem],
                                                      'Removed': [doc.Models[0].RootItem]}, doc)
        finally:
            sys.modules.pop('pynavis.sets', None)
    assert written == ['Changed'] and skipped == ['Added'] and failed == ['Removed']


def test_snapshot_store_round_trip():
    from namo import paths
    folder = _compat.tmpdir('modelcompare_store' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        doc = _versions()
        with nav.FakeNav(doc):
            sys.modules['pynavis.doc'].get_filename = lambda d=None: os.path.join(folder, 'Site.nwf')
            assert mc.list_snapshots(doc) == []
            data = mc.collect(doc.Models[0].RootItem)
            snap = mc.snapshot_record('Before', 'V1.nwc', 'Meters', data['records'], [], '2026-10-05 10:00')
            stored, note = mc.save_snapshot(snap, doc)
            later = mc.snapshot_record('After', 'V1.nwc', 'Meters', data['records'], [], '2026-10-05 11:00')
            mc.save_snapshot(later, doc)
            listed = mc.list_snapshots(doc)
            loaded = mc.load_snapshot(stored['id'], doc)
    finally:
        paths.redirect(previous)
    assert stored['id'] == 'S-001' and note is None
    assert [s['label'] for s in listed] == ['After', 'Before'] and 'objects' not in listed[0]
    assert mc.check_snapshot(loaded) is None and len(loaded['objects']) == 4
    assert [r['name'] for r in mc.decode_records(loaded['objects'])][0] == 'Kept'
