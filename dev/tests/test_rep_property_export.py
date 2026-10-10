# -*- coding: utf-8 -*-
"""rep_property_export: contract, headers, cells and xlsx/csv runs through the fake layer."""
import io
import os

import _compat
import test_propindex as pf

from namo import propindex, reports, rep_property_export as rep


def test_contract_and_headers():
    config = rep.default_config()
    assert 'folder' in config and 'name_template' in config
    assert rep.TITLE == reports.title_of('property_export')
    assert rep.describe(dict(config, properties=[['A', 'B']])) == '1 property, current selection, XLSX'
    headers = rep.column_headers(['ifc_guid', 'name'], [('Element', 'Length'), ('Element', 'Mark')], {0: 'mm'})
    assert headers == ['Name', 'IFC GlobalId', 'Element > Length (mm)', 'Element > Mark']


def test_cells():
    ctx = propindex.context_for('Millimeters')
    assert rep.cell_for('bool', False, ctx) == 'No'
    assert _compat.approx(rep.cell_for('length', 1.5, ctx), 1500.0)
    assert rep.cell_for('int', 7, ctx) == 7
    assert rep.cell_for('text', '  ', ctx) is None
    assert rep.cell_for('text', None, ctx) is None


def test_run_xlsx_and_csv():
    folder = _compat.tmpdir('g2_rep_export')
    doc, n = pf.tree()
    with pf.Fake(doc, 'Millimeters'):
        doc.CurrentSelection.SelectedItems = [n['Level 1'], n['Duct']]
        config = dict(rep.default_config(), properties=[['Element', 'Length'], ['Element', 'Mark']],
                      path=os.path.join(folder, 'p.xlsx'))
        result = rep.run(config, doc=doc, interactive=False)
        assert result['level'] == 'success' and result['message'] == 'Exported 3 rows', result
        csv_config = dict(config, format='csv', path=os.path.join(folder, 'p.csv'), identity=['name'])
        result = rep.run(csv_config, doc=doc, interactive=False)
        assert result['paths'] == [os.path.join(folder, 'p.csv')]
        nothing = rep.run(dict(config, properties=[], identity=[]), doc=doc, interactive=False)
        assert nothing['level'] == 'warning'
        doc.CurrentSelection.SelectedItems = []
        assert rep.run(config, doc=doc, interactive=False)['message'] == 'Nothing selected'
    with io.open(os.path.join(folder, 'p.csv'), encoding='utf-8-sig') as handle:
        lines = handle.read().splitlines()
    assert lines[0] == 'Name,Element > Length (mm),Element > Mark'
    assert lines[1] == 'Wall 1,3000.0,W1'
    assert lines[3] == 'Duct,2000.0,'
    try:
        import openpyxl
    except ImportError:
        return
    wb = openpyxl.load_workbook(os.path.join(folder, 'p.xlsx'))
    assert wb.sheetnames == ['Properties', 'About']
    ws = wb['Properties']
    assert [c.value for c in ws[1]] == ['Name', 'Source file', 'Element ID', 'IFC GlobalId',
                                        'Element > Length (mm)', 'Element > Mark']
    assert ws.cell(2, 1).value == 'Wall 1' and ws.cell(2, 5).value == 3000
    assert ws.cell(2, 3).value == '101'
