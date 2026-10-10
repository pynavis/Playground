# -*- coding: utf-8 -*-
"""namo.xlsx: package structure and OOXML element order (both engines), openpyxl
round trips of values, styles, merges, freeze panes, autofilter, links, images (CPython)."""
import datetime
import io
import os
import re
import zipfile
from xml.dom import minidom

import _compat
from _compat import raises, tmpdir

from namo import png, xlsx

CREATED = datetime.datetime(2026, 10, 5, 9, 30, 0)

# CT_Worksheet / CT_Workbook / CT_Stylesheet child sequences from the OOXML schema.
WORKSHEET_ORDER = [
    'sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols', 'sheetData', 'sheetCalcPr',
    'sheetProtection', 'protectedRanges', 'scenarios', 'autoFilter', 'sortState',
    'dataConsolidate', 'customSheetViews', 'mergeCells', 'phoneticPr', 'conditionalFormatting',
    'dataValidations', 'hyperlinks', 'printOptions', 'pageMargins', 'pageSetup', 'headerFooter',
    'rowBreaks', 'colBreaks', 'customProperties', 'cellWatches', 'ignoredErrors', 'smartTags',
    'drawing', 'legacyDrawing', 'legacyDrawingHF', 'picture', 'oleObjects', 'controls',
    'webPublishItems', 'tableParts', 'extLst']
WORKBOOK_ORDER = ['fileVersion', 'fileSharing', 'workbookPr', 'workbookProtection', 'bookViews',
                  'sheets', 'functionGroups', 'externalReferences', 'definedNames', 'calcPr']
STYLES_ORDER = ['numFmts', 'fonts', 'fills', 'borders', 'cellStyleXfs', 'cellXfs', 'cellStyles',
                'dxfs', 'tableStyles', 'colors', 'extLst']
BORDER_ORDER = ['start', 'left', 'end', 'right', 'top', 'bottom', 'diagonal', 'vertical', 'horizontal']


def _thumb(width=48, height=27, rgb=(30, 90, 160)):
    return png.encode(width, height, list(rgb) * (width * height), 'RGB')


def _sample():
    wb = xlsx.Workbook(title='Clash report & more', author='Tester', created=CREATED)
    ws = wb.sheet('Clashes')
    ws.title('Level 3 <coordination>', span=5)
    ws.subtitle('Generated for testing', span=5)
    ws.columns([('Name', 28, None), ('Status', 12, None), ('Distance', None, 'num2'),
                ('Found', 12, None), ('Image', 10, None)])
    ws.row(['Clash1 & <b>', 'New', -0.0425, datetime.date(2026, 10, 1)],
           styles=[None, xlsx.status_style('New')])
    ws.row(['Clash2', 'Resolved', 1234.5, datetime.datetime(2026, 10, 2, 14, 30, 15)],
           styles=[None, xlsx.status_style('resolved')])
    ws.row([xlsx.Link('https://example.com/a b?x=1&y=2', 'Web link', tip='Opens a browser'),
            True, 7, None, None])
    ws.row([xlsx.Link(r'C:\Temp\my report #1.pdf', 'File link'), False, 0.1, None])
    ws.row([xlsx.Link('#' + xlsx.quote_sheet('Second sheet') + '!B2'), None, float('nan')])
    ws.write(8, 0, xlsx.Formula('=SUM(C4:C7)'))
    ws.write(8, 1, xlsx.Formula('C4*2', value=-0.085))
    ws.image(_thumb(), 3, 4, width_px=96, fit_cell=True, description='Thumb')
    ws.image(_thumb(), 4, 4, height_px=20)
    ws.autofilter()
    ws.tab_color('2A78D6')
    second = wb.sheet('Second sheet')
    second.columns(['Model', 'Count'])
    for i in range(5):
        second.row(['Model %d' % i, i * 1000], styles=[None, 'int'])
    second.write(1, 3, 5, wb.fill_style('184F95', base='int'))
    second.write(2, 3, 2, wb.fill_style('CDE2FB', base='int'))
    second.merge(7, 0, 8, 2)
    return wb


def _parts(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    try:
        return dict((name, archive.read(name)) for name in archive.namelist()), archive.namelist()
    finally:
        archive.close()


def _children(xml_bytes):
    root = minidom.parseString(xml_bytes).documentElement
    return root, [n.localName or n.tagName for n in root.childNodes if n.nodeType == n.ELEMENT_NODE]


def _assert_order(names, order):
    positions = [order.index(n) for n in names]
    assert positions == sorted(positions), names
    assert len(set(names)) == len(names), names


# --- pure helpers ----------------------------------------------------------------------

def test_col_letters_and_refs():
    assert [xlsx.col_letter(i) for i in (0, 25, 26, 51, 52, 701, 702, 16383)] == \
        ['A', 'Z', 'AA', 'AZ', 'BA', 'ZZ', 'AAA', 'XFD']
    assert xlsx.cell_ref(0, 0) == 'A1'
    assert xlsx.cell_ref(9, 2, absolute=True) == '$C$10'
    assert xlsx.range_ref(4, 3, 1, 0) == 'A2:D5'
    assert xlsx.range_ref(2, 2, 2, 2) == 'C3'
    raises(ValueError, xlsx.col_letter, -1)


def test_sheet_name_sanitising():
    assert xlsx.sanitize_sheet_name('a[b]:c*d?e/f\\g') == 'a b c d e f g'
    assert xlsx.sanitize_sheet_name('') == 'Sheet'
    assert xlsx.sanitize_sheet_name(None) == 'Sheet'
    assert xlsx.sanitize_sheet_name("'quoted'") == 'quoted'
    assert xlsx.sanitize_sheet_name('History') == 'History sheet'
    long_name = 'x' * 40
    assert xlsx.sanitize_sheet_name(long_name) == 'x' * 31
    assert xlsx.sanitize_sheet_name('Data', ['data']) == 'Data (2)'
    assert xlsx.sanitize_sheet_name('Data', ['DATA', 'Data (2)']) == 'Data (3)'
    deduped = xlsx.sanitize_sheet_name(long_name, ['x' * 31])
    assert deduped == 'x' * 27 + ' (2)' and len(deduped) == 31


def test_workbook_dedupes_sheet_names():
    wb = xlsx.Workbook(created=CREATED)
    names = [wb.sheet(n).name for n in ('Clashes', 'clashes', 'A/B', 'Clashes')]
    assert names == ['Clashes', 'clashes (2)', 'A B', 'Clashes (3)']
    assert wb.get_sheet('CLASHES').name == 'Clashes'
    assert wb.get_sheet('nope') is None


def test_number_text_is_shortest_and_engine_independent():
    assert xlsx._num(0.1) == '0.1'
    assert xlsx._num(1.0 / 3) == '0.3333333333333333'
    assert xlsx._num(1e16) == '1e+16'
    assert xlsx._num(-0.0425) == '-0.0425'
    assert xlsx._num(12) == '12'
    assert xlsx._num(10 ** 20) == '100000000000000000000'
    assert xlsx._num(True) == '1'
    for value in (0.1, 2.675, 1.5e-7, 123456789.123, 1.0 / 7):
        assert float(xlsx._num(value)) == value


def test_excel_serial_dates():
    assert xlsx.excel_serial(datetime.date(1900, 3, 1)) == 61
    assert xlsx.excel_serial(datetime.date(2026, 10, 5)) == 46300
    assert xlsx.excel_serial(datetime.datetime(2026, 10, 5, 12, 0)) == 46300.5
    assert xlsx.excel_serial(datetime.time(6, 0)) == 0.25
    raises(TypeError, xlsx.excel_serial, '2026-10-05')


def test_link_targets():
    assert xlsx._link_target('https://x.com/a b') == ('https://x.com/a%20b', None)
    assert xlsx._link_target('mailto:a@b.com') == ('mailto:a@b.com', None)
    assert xlsx._link_target(r'C:\A dir\f#1 50%.pdf') == ('file:///C:/A%20dir/f%231%2050%25.pdf', None)
    assert xlsx._link_target(r'\\server\share\x.nwd') == ('file://server/share/x.nwd', None)
    assert xlsx._link_target(r'images\a.png') == ('images/a.png', None)
    assert xlsx._link_target("#'My sheet'!A1") == (None, "'My sheet'!A1")
    assert xlsx.Link('#Sheet2!A1').text == 'Sheet2!A1'
    assert xlsx.Link('https://x.com').text == 'https://x.com'


def test_status_styles_and_formula():
    assert xlsx.status_style('New') == 'status_new'
    assert xlsx.status_style(' APPROVED ') == 'status_approved'
    assert xlsx.status_style('Unknown') is None
    assert xlsx.status_style(None) is None
    assert xlsx.Formula('=A1+1').text == 'A1+1'
    assert xlsx.Formula('A1').text == 'A1'


def test_style_definitions_and_validation():
    wb = xlsx.Workbook(created=CREATED)
    for name in ('header', 'title', 'muted', 'int', 'num1', 'num2', 'pct', 'date', 'datetime',
                 'good', 'warn', 'bad', 'info', 'link', 'wrap', 'status_new', 'status_active',
                 'status_reviewed', 'status_approved', 'status_resolved'):
        assert wb.has_style(name), name
    assert wb.style('mine', bold=True, fill='#ffcc00', border='top bottom', align='right') == 'mine'
    assert wb._styles['mine']['fill'] == 'FFFFCC00'
    assert wb._styles['mine']['border'] == ('top', 'bottom')
    wb.style('derived', base='mine', bold=False, color=(255, 0, 0))
    derived = wb._styles['derived']
    assert derived['fill'] == 'FFFFCC00' and not derived['bold'] and derived['color'] == 'FFFF0000'
    assert derived['border'] == ('top', 'bottom')
    raises(ValueError, wb.style, 'bad1', fill='red')
    raises(ValueError, wb.style, 'bad2', align='middle-ish')
    raises(ValueError, wb.style, 'bad3', border='diagonal')
    raises(ValueError, wb.style, 'bad4', base='missing')
    ws = wb.sheet('S')
    raises(ValueError, ws.write, 0, 0, 1, 'no-such-style')
    raises(ValueError, ws.write, xlsx.MAX_ROWS, 0, 1)
    raises(ValueError, ws.write, 0, xlsx.MAX_COLS, 1)


def test_fill_style_is_cached_and_readable():
    wb = xlsx.Workbook(created=CREATED)
    dark = wb.fill_style('0D366B', base='int')
    assert dark == wb.fill_style('#0d366b', base='int')
    assert wb._styles[dark]['color'] == 'FFFFFFFF'
    assert wb._styles[dark]['num_format'] == '#,##0'
    light = wb.fill_style('CDE2FB')
    assert wb._styles[light]['color'] is None


def test_merge_rules():
    ws = xlsx.Workbook(created=CREATED).sheet('S')
    ws.merge(0, 0, 0, 3)
    ws.merge(2, 2, 2, 2)                     # single cell: ignored
    raises(ValueError, ws.merge, 0, 2, 1, 4)  # overlaps A1:D1
    ws.merge(1, 0, 1, 3)
    assert ws._merges == [(0, 0, 0, 3), (1, 0, 1, 3)]


def test_append_positions_and_column_defaults():
    ws = xlsx.Workbook(created=CREATED).sheet('S')
    assert ws.title('T', span=3) == 0
    assert ws.columns([('A', 10, 'int'), ('B', None, None)]) == 1
    assert ws._freeze == (2, 0)
    row = ws.row([5, datetime.date(2026, 1, 1)])
    assert row == 2 and ws.cell(2, 0) == (5, 'int') and ws.cell(2, 1)[1] == 'date'
    ws.write(10, 0, 'far')
    assert ws.next_row == 11
    assert ws.row(['next']) == 11
    ws.skip(2)
    assert ws.next_row == 14
    assert ws.used_range() == (0, 0, 11, 1)


# --- package structure (both engines) ---------------------------------------------------

def test_package_parts_and_relationships():
    data = _sample().to_bytes()
    parts, order = _parts(data)
    assert order[0] == '[Content_Types].xml'
    expected = ['_rels/.rels', 'docProps/core.xml', 'docProps/app.xml', 'xl/workbook.xml',
                'xl/_rels/workbook.xml.rels', 'xl/styles.xml', 'xl/sharedStrings.xml',
                'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml',
                'xl/worksheets/_rels/sheet1.xml.rels', 'xl/drawings/drawing1.xml',
                'xl/drawings/_rels/drawing1.xml.rels', 'xl/media/image1.png']
    for name in expected:
        assert name in parts, name
    # identical image bytes are stored once
    assert len([n for n in parts if n.startswith('xl/media/')]) == 1
    for name, content in parts.items():
        if name.endswith('.xml') or name.endswith('.rels'):
            minidom.parseString(content)          # well-formed
    types = parts['[Content_Types].xml'].decode('utf-8')
    for name in parts:
        if name.endswith('.xml') and name != '[Content_Types].xml' and not name.startswith('_rels'):
            assert ('PartName="/%s"' % name) in types, name
    assert 'Extension="png" ContentType="image/png"' in types
    # every internal relationship target exists
    for name, content in parts.items():
        if not name.endswith('.rels'):
            continue
        base = name.replace('_rels/', '').replace('.rels', '')
        folder = os.path.dirname(base)
        for target, mode in re.findall(r'Target="([^"]+)"( TargetMode="External")?',
                                       content.decode('utf-8')):
            if mode:
                continue
            resolved = os.path.normpath(os.path.join(folder, target)).replace('\\', '/')
            assert resolved in parts, (name, target)


def test_worksheet_element_order_and_cells():
    parts, _ = _parts(_sample().to_bytes())
    for name in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
        root, names = _children(parts[name])
        _assert_order(names, WORKSHEET_ORDER)
        rows = root.getElementsByTagName('row')
        numbers = [int(r.getAttribute('r')) for r in rows]
        assert numbers == sorted(numbers) and len(set(numbers)) == len(numbers)
        for r in rows:
            refs = [c.getAttribute('r') for c in r.getElementsByTagName('c')]
            cols = []
            for ref in refs:
                match = re.match(r'^([A-Z]{1,3})([0-9]+)$', ref)
                assert match and int(match.group(2)) == int(r.getAttribute('r')), ref
                number = 0
                for ch in match.group(1):
                    number = number * 26 + ord(ch) - 64
                cols.append(number)
            assert cols == sorted(cols)
    _root, names = _children(parts['xl/worksheets/sheet1.xml'])
    assert names == ['sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols', 'sheetData',
                     'autoFilter', 'mergeCells', 'hyperlinks', 'pageMargins', 'drawing']


def test_workbook_and_styles_parts():
    parts, _ = _parts(_sample().to_bytes())
    _root, names = _children(parts['xl/workbook.xml'])
    _assert_order(names, WORKBOOK_ORDER)
    text = parts['xl/workbook.xml'].decode('utf-8')
    assert 'fullCalcOnLoad="1"' in text
    assert "_xlnm._FilterDatabase" in text and "'Clashes'!$A$3:$E$9" in text
    root, names = _children(parts['xl/styles.xml'])
    _assert_order(names, STYLES_ORDER)
    for fmt in root.getElementsByTagName('numFmt'):
        assert int(fmt.getAttribute('numFmtId')) >= 164
    fills = root.getElementsByTagName('fill')
    assert fills[0].getElementsByTagName('patternFill')[0].getAttribute('patternType') == 'none'
    assert fills[1].getElementsByTagName('patternFill')[0].getAttribute('patternType') == 'gray125'
    for tag in ('fonts', 'fills', 'borders', 'cellXfs', 'numFmts'):
        node = root.getElementsByTagName(tag)[0]
        kids = [n for n in node.childNodes if n.nodeType == n.ELEMENT_NODE]
        assert int(node.getAttribute('count')) == len(kids), tag
    xf0 = root.getElementsByTagName('cellXfs')[0].getElementsByTagName('xf')[0]
    assert [xf0.getAttribute(a) for a in ('numFmtId', 'fontId', 'fillId', 'borderId')] == ['0'] * 4
    for border in root.getElementsByTagName('border'):
        kids = [n.tagName for n in border.childNodes if n.nodeType == n.ELEMENT_NODE]
        _assert_order(kids, BORDER_ORDER)
    xf_count = len(root.getElementsByTagName('cellXfs')[0].getElementsByTagName('xf'))
    for sheet in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
        for index in re.findall(r' s="([0-9]+)"', parts[sheet].decode('utf-8')):
            assert int(index) < xf_count


def test_text_escaping_and_shared_strings():
    wb = xlsx.Workbook(created=CREATED)
    ws = wb.sheet('S')
    ws.row([u'a & b < c > "d"', u'ctrl\x01\x0bchars', u'line1\r\nline2', u'_x0041_', u'same', u'same'])
    ws.write(1, 0, 'z' * 40000)
    parts, _ = _parts(wb.to_bytes())
    sst = parts['xl/sharedStrings.xml'].decode('utf-8')
    assert 'a &amp; b &lt; c &gt; "d"' in sst
    assert 'ctrlchars' in sst and '\x01' not in sst
    assert 'line1\nline2' in sst
    assert '_x005F_x0041_' in sst
    assert sst.count('>same<') == 1
    assert 'count="7" uniqueCount="6"' in sst
    assert ('z' * xlsx.MAX_TEXT + '<') in sst and ('z' * (xlsx.MAX_TEXT + 1)) not in sst


def test_output_is_deterministic():
    assert _sample().to_bytes() == _sample().to_bytes()


def test_empty_workbook_still_has_a_sheet():
    parts, _ = _parts(xlsx.Workbook(created=CREATED).to_bytes())
    assert 'xl/worksheets/sheet1.xml' in parts
    assert '<sheet name="Sheet1"' in parts['xl/workbook.xml'].decode('utf-8')


def test_save_is_atomic_and_returns_path():
    folder = tmpdir('xlsx_save')
    path = os.path.join(folder, 'sub', 'out.xlsx')
    assert _sample().save(path) == path
    assert os.path.exists(path) and not os.path.exists(path + '.tmp')
    assert zipfile.ZipFile(path).testzip() is None


def test_write_simple():
    folder = tmpdir('xlsx_simple')
    path = os.path.join(folder, 'simple.xlsx')
    sheets = [{'name': 'Counts', 'headers': ['Category', 'Count'], 'rows': [['Ducts', 12], ['Pipes', 3]],
               'widths': [20, None], 'styles': [None, 'int'], 'title': 'Totals'},
              {'name': 'Raw', 'rows': [['no headers', 1]]}]
    assert xlsx.write_simple(path, sheets) == path
    parts, _ = _parts(open(path, 'rb').read())
    sheet1 = parts['xl/worksheets/sheet1.xml'].decode('utf-8')
    assert '<autoFilter ref="A2:B4"/>' in sheet1 and 'ySplit="2"' in sheet1
    assert '<mergeCell ref="A1:B1"/>' in sheet1


def test_images_scale_and_reject_unknown_types():
    ws = xlsx.Workbook(created=CREATED).sheet('S')
    assert ws.image(_thumb(100, 50), 0, 0) == (100, 50)
    assert ws.image(_thumb(100, 50), 1, 0, width_px=50) == (50, 25)
    assert ws.image(_thumb(100, 50), 2, 0, height_px=100) == (200, 100)
    assert ws.image(_thumb(100, 50), 3, 0, width_px=10, height_px=10) == (10, 10)
    ws.image(_thumb(100, 60), 4, 1, fit_cell=True)
    assert ws._heights[4] == xlsx.px_to_points(62)
    assert abs(ws._widths[1] - xlsx.px_to_chars(102)) < 1e-9
    raises(ValueError, ws.image, b'not an image', 0, 0)
    raises(ValueError, ws.image, _thumb(), 0, 0, anchor='float')


def test_two_cell_anchor_ends_in_the_right_cell():
    wb = xlsx.Workbook(created=CREATED)
    ws = wb.sheet('S')
    ws.set_width(1, 10)                       # 75 px
    ws.set_height(2, 30)                      # 40 px
    ws.image(_thumb(100, 50), 2, 1, offset_px=(4, 4), anchor='cell')       # crosses into C and row 4
    ws.image(_thumb(20, 10), 5, 1, fit_cell=True, anchor='cell')           # stays inside B6
    parts, _ = _parts(wb.to_bytes())
    drawing = minidom.parseString(parts['xl/drawings/drawing1.xml'])
    anchors = drawing.getElementsByTagName('xdr:twoCellAnchor')
    assert len(anchors) == 2 and anchors[0].getAttribute('editAs') == 'twoCell'

    def corner(anchor, tag):
        node = anchor.getElementsByTagName(tag)[0]
        return tuple(int(node.getElementsByTagName('xdr:' + k)[0].firstChild.data)
                     for k in ('col', 'colOff', 'row', 'rowOff'))

    assert corner(anchors[0], 'xdr:from') == (1, 4 * xlsx.EMU_PER_PX, 2, 4 * xlsx.EMU_PER_PX)
    # 104 px from the left of B (75 px wide) ends 29 px into C (default 64 px);
    # 54 px from the top of row 3 (40 px) ends 14 px into row 4 (default 20 px)
    assert corner(anchors[0], 'xdr:to') == (2, 29 * xlsx.EMU_PER_PX, 3, 14 * xlsx.EMU_PER_PX)
    assert corner(anchors[1], 'xdr:to') == (1, 20 * xlsx.EMU_PER_PX, 5, 10 * xlsx.EMU_PER_PX)


# --- openpyxl round trip (CPython) ------------------------------------------------------

def _openpyxl():
    _compat.require_cpython('openpyxl runs under CPython')
    try:
        import openpyxl
    except ImportError:
        raise _compat.SkipTest('openpyxl is not installed')
    return openpyxl


def _load(wb):
    openpyxl = _openpyxl()
    import warnings
    folder = tmpdir('xlsx_openpyxl')
    path = wb.save(os.path.join(folder, 'book.xlsx'))
    with warnings.catch_warnings():
        warnings.simplefilter('error')       # any openpyxl complaint fails the test
        return openpyxl.load_workbook(path)


def test_openpyxl_values_and_types():
    book = _load(_sample())
    assert book.sheetnames == ['Clashes', 'Second sheet']
    ws = book['Clashes']
    assert ws['A1'].value == 'Level 3 <coordination>'
    assert [c.value for c in ws[3]] == ['Name', 'Status', 'Distance', 'Found', 'Image']
    assert ws['A4'].value == 'Clash1 & <b>'
    assert ws['C4'].value == -0.0425
    assert ws['D4'].value == datetime.datetime(2026, 10, 1)
    assert ws['D5'].value == datetime.datetime(2026, 10, 2, 14, 30, 15)
    assert ws['B6'].value is True and ws['B7'].value is False
    assert ws['C6'].value == 7 and ws['C7'].value == 0.1
    assert ws['C8'].value == '#NUM!'
    assert ws['A9'].value == '=SUM(C4:C7)' and ws['B9'].value == '=C4*2'
    assert ws['A6'].value == 'Web link'
    assert book.properties.title == 'Clash report & more'
    assert book.properties.creator == 'Tester'


def test_openpyxl_styles():
    book = _load(_sample())
    ws = book['Clashes']
    header = ws['A3']
    assert header.font.b and header.fill.fgColor.rgb == 'FFF2F2F2'
    assert header.border.bottom.style == 'thin' and header.border.bottom.color.rgb == 'FFA6A6A6'
    assert ws['A1'].font.b and ws['A1'].font.sz == 14
    assert ws['A2'].font.color.rgb == 'FF6B6B6B'
    assert ws['B4'].fill.fgColor.rgb == 'FFF9D3D3' and ws['B4'].font.color.rgb == 'FF8E1414'
    assert ws['B5'].fill.fgColor.rgb == 'FFFFF1B8'
    assert ws['C4'].number_format == '#,##0.00'
    assert ws['D4'].number_format == 'yyyy-mm-dd'
    assert ws['D5'].number_format == 'yyyy-mm-dd hh:mm'
    assert ws['A6'].font.u == 'single' and ws['A6'].font.color.rgb == 'FF0563C1'
    assert ws.sheet_properties.tabColor.rgb == 'FF2A78D6'
    second = book['Second sheet']
    assert second['B3'].number_format == '#,##0'
    assert second['D2'].fill.fgColor.rgb == 'FF184F95' and second['D2'].font.color.rgb == 'FFFFFFFF'
    assert second['D3'].fill.fgColor.rgb == 'FFCDE2FB'


def test_openpyxl_layout_merges_freeze_filter():
    book = _load(_sample())
    ws = book['Clashes']
    assert sorted(str(r) for r in ws.merged_cells.ranges) == ['A1:E1', 'A2:E2']
    assert ws.freeze_panes == 'A4'
    assert ws.auto_filter.ref == 'A3:E9'
    assert ws.column_dimensions['A'].width == 28
    assert ws.column_dimensions['C'].width >= 10        # auto width from content
    assert ws.row_dimensions[1].height == 22
    second = book['Second sheet']
    assert second.freeze_panes == 'A2'
    assert [str(r) for r in second.merged_cells.ranges] == ['A8:C9']
    assert second.auto_filter.ref is None


def test_openpyxl_hyperlinks_and_images():
    book = _load(_sample())
    ws = book['Clashes']
    assert ws['A6'].hyperlink.target == 'https://example.com/a%20b?x=1&y=2'
    assert ws['A6'].hyperlink.tooltip == 'Opens a browser'
    assert ws['A7'].hyperlink.target == 'file:///C:/Temp/my%20report%20%231.pdf'
    assert ws['A8'].hyperlink.location == "'Second sheet'!B2"
    assert len(ws._images) == 2
    anchor = ws._images[0].anchor
    assert anchor._from.row == 3 and anchor._from.col == 4
    assert anchor.ext.width == 96 * xlsx.EMU_PER_PX and anchor.ext.height == 54 * xlsx.EMU_PER_PX
    assert len(book['Second sheet']._images) == 0


def test_openpyxl_reads_two_cell_anchors():
    wb = xlsx.Workbook(created=CREATED)
    ws = wb.sheet('S')
    ws.columns([('Name', 12, None), ('Image', None, None)])
    for i in range(3):
        r = ws.row(['Row %d' % i])
        ws.image(_thumb(64, 36), r, 1, offset_px=(2, 2), fit_cell=True, anchor='cell')
    ws.autofilter()
    book = _load(wb)
    images = book['S']._images
    assert len(images) == 3
    anchor = images[2].anchor
    assert anchor._from.row == 3 and anchor._from.col == 1 and anchor.to.row == 3 and anchor.to.col == 1


def test_openpyxl_jpeg_and_unicode():
    _compat.require_cpython('Pillow runs under CPython')
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (64, 32), (200, 40, 40)).save(buffer, 'JPEG')
    wb = xlsx.Workbook(created=CREATED)
    ws = wb.sheet(u'Unicode \u00e9\u00e8 \u4e2d\u6587')
    ws.row([u'caf\u00e9', u'\u4e2d\u6587', u'emoji \U0001F600', u'tab\there'])
    ws.image(buffer.getvalue(), 2, 0)
    book = _load(wb)
    sheet = book.worksheets[0]
    assert sheet.title == u'Unicode \u00e9\u00e8 \u4e2d\u6587'
    assert [c.value for c in sheet[1]] == [u'caf\u00e9', u'\u4e2d\u6587', u'emoji \U0001F600', u'tab\there']
    assert len(sheet._images) == 1
    parts, _ = _parts(wb.to_bytes())
    assert 'xl/media/image1.jpeg' in parts
    assert 'Extension="jpeg" ContentType="image/jpeg"' in parts['[Content_Types].xml'].decode('utf-8')


def test_openpyxl_large_sheet():
    wb = xlsx.Workbook(created=CREATED)
    ws = wb.sheet('Big')
    ws.columns([('Id', 8, 'int'), ('Name', 20, None), ('Value', 10, 'num2')])
    for i in range(3000):
        ws.row([i, 'Item %d' % (i % 50), i * 0.5])
    ws.autofilter()
    book = _load(wb)
    sheet = book['Big']
    assert sheet.max_row == 3001
    assert sheet['B3001'].value == 'Item %d' % (2999 % 50)
    assert sheet.auto_filter.ref == 'A1:C3001'


# --- same bytes on both engines ---------------------------------------------------------

GOLDEN_XLSX = 'bac01c006235ea6e7a8ae4f4bda3ce25'


def test_golden_digest_matches_on_both_engines():
    """Every uncompressed part hashes the same under CPython and IronPython (floats, dict
    order, dates). If you change the output on purpose, update GOLDEN_XLSX after
    checking both engines agree."""
    import hashlib
    wb = xlsx.Workbook(title='Golden', author='Both engines', created=CREATED)
    ws = wb.sheet('Data')
    ws.title('Golden workbook', span=4)
    ws.columns([('Name', 20, None), ('Value', None, 'num2'), ('When', 12, None), ('Status', 10, None)])
    for i in range(12):
        status = ['New', 'Active', 'Resolved'][i % 3]
        ws.row([u'Item %d \u00e9' % i, i / 7.0, datetime.datetime(2026, 1, 1 + i, i, 30), status],
               styles=[None, None, None, xlsx.status_style(status)])
    ws.write(14, 1, xlsx.Formula('SUM(B3:B14)', value=1.0 / 3))
    ws.write(15, 0, xlsx.Link('C:\\Temp\\a b.pdf', 'File'))
    ws.write(16, 0, 1e-7, wb.fill_style('3987E5', base='num2'))
    ws.autofilter()
    archive = zipfile.ZipFile(io.BytesIO(wb.to_bytes()))
    digest = hashlib.md5()
    for name in archive.namelist():
        digest.update(name.encode('utf-8'))
        digest.update(archive.read(name))
    assert digest.hexdigest() == GOLDEN_XLSX
