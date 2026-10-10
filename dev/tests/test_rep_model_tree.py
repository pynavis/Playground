# -*- coding: utf-8 -*-
"""rep_model_tree: bottom-up counts, rows, collect() and a run through the fake layer."""
import os

import _compat
import test_propindex as pf

from namo import reports, rep_model_tree as rep


def test_contract():
    config = rep.default_config()
    assert 'folder' in config and 'name_template' in config and config['depth'] == 4
    assert rep.TITLE == reports.title_of('model_tree')
    assert rep.describe(config) == 'Depth 4, every model'


def test_aggregate_counts_subtrees():
    nodes = [{'depth': 0, 'own_geometry': False}, {'depth': 1, 'own_geometry': False},
             {'depth': 2, 'own_geometry': True}, {'depth': 2, 'own_geometry': False, 'below': (5, 3)},
             {'depth': 1, 'own_geometry': True}]
    rep.aggregate(nodes)
    assert [(n['descendants'], n['geometry']) for n in nodes] == [(9, 5), (7, 4), (0, 1), (5, 3), (0, 1)]


def test_rows_and_headers():
    nodes = [{'depth': 0, 'name': 'A.nwc', 'class': 'File', 'descendants': 2, 'geometry': 1, 'hidden': False},
             {'depth': 1, 'name': 'L1', 'class': 'Layer', 'descendants': 0, 'geometry': 0, 'hidden': True}]
    rows = rep.sheet_rows(nodes, 1)
    assert rows == [[0, 'A.nwc', '', 'A.nwc', 'File', 2, 1, ''], [1, '', 'L1', 'L1', 'Layer', 0, 0, 'Yes']]
    assert rep.headers(1) == ['Depth', 'Level 0', 'Level 1', 'Name', 'Class', 'Descendants', 'Geometry items',
                              'Hidden']


def test_collect_and_run():
    folder = _compat.tmpdir('g2_rep_tree')
    doc, n = pf.tree()
    with pf.Fake(doc):
        from namo import items
        models = items.models(doc)
        nodes, truncated = rep.collect(models[0], 1)
        assert not truncated
        assert [(x['name'], x['depth'], x['descendants'], x['geometry']) for x in nodes] == [
            ('ARCH.nwc', 0, 5, 2), ('Level 1', 1, 4, 2)]
        nodes, _t = rep.collect(models[1], 4)
        assert [(x['name'], x['hidden']) for x in nodes] == [('MEP.nwc', False), ('Duct', False),
                                                              ('Duct solid', False), ('Pipe', True)]
        capped, truncated = rep.collect(models[0], 4, max_rows=2)
        assert truncated and len(capped) == 2
        path = os.path.join(folder, 'tree.xlsx')
        result = rep.run(dict(rep.default_config(), path=path, depth=2), doc=doc, interactive=False)
        assert result['level'] == 'success', result
        assert result['message'] == 'Exported 8 tree nodes from 2 models'
        chosen = rep.run(dict(rep.default_config(), path=path, models=['nope.rvt']), doc=doc, interactive=False)
        assert chosen['level'] == 'warning'
    try:
        import openpyxl
    except ImportError:
        return
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Summary', 'ARCH.rvt', 'MEP.rvt']
    ws = wb['ARCH.rvt']
    assert [c.value for c in ws[1]][:4] == ['Depth', 'Level 0', 'Level 1', 'Level 2']
    assert ws.cell(3, 3).value == 'Level 1' and ws.cell(3, 5).value == 'Level 1'
    assert ws.cell(3, 5).alignment.indent == 1
