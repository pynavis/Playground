# -*- coding: utf-8 -*-
"""rep_clash_excel: the workbook built from realistic fake clashes, read back with openpyxl
(sheets, fills, filters, frozen panes, pictures, links), and run() through the fake
Navisworks layer."""
import os

import _compat

import test_clashreport as fx
from namo import clashreport as cr
from namo import png, reports
from namo import rep_clash_excel as rep


def _thumbs(data, folder, per_test=5):
    """Writes small PNG thumbnails for the first listed rows of each test."""
    pixels = bytes(bytearray([200, 120, 40] * (cr.THUMB_W * cr.THUMB_H)))
    for t_index, test in enumerate(data['tests']):
        for row in cr.filter_rows(test['rows'], cr.OPEN)[:per_test]:
            path = os.path.join(folder, 't%d-%d.png' % (t_index, row['id']))
            png.write(path, cr.THUMB_W, cr.THUMB_H, pixels, 'RGB')
            row['image'] = path


def _openpyxl():
    _compat.require_cpython('openpyxl reads the workbook under CPython')
    import openpyxl
    return openpyxl


def test_contract():
    config = rep.default_config()
    assert config['folder'] == '' and config['name_template'] == rep.NAME_TEMPLATE
    assert config['statuses'] == cr.OPEN and config['images'] is True
    assert reports.title_of('clash_excel') == rep.TITLE
    assert rep.describe(config) == 'Every test with results, open clashes (New, Active, Reviewed), with thumbnails'
    assert rep.describe({'tests': ['A'], 'statuses': [], 'images': False}) == 'A, all statuses'


def test_columns_values_and_styles_line_up():
    for images in (True, False):
        for with_test in (True, False):
            cols = rep.clash_columns('mm', images, with_test)
            row = fx.fake_data()['tests'][0]['rows'][0]
            values = rep.clash_values(row, images, with_test)
            assert len(values) == len(cols)
            styles = rep.clash_styles(row, images, with_test)
            status_col = [h for h, _w, _s in cols].index('Status')
            assert styles[status_col] == 'c_' + row['status'].lower()
            assert values[status_col] == row['status']
    assert rep.clash_columns('ft', False)[3][0] == 'Distance (ft)'


def test_workbook_structure():
    openpyxl = _openpyxl()
    import warnings
    folder = _compat.tmpdir('rep-clash-excel-wb')
    data = fx.fake_data()
    _thumbs(data, folder)
    data['notes'].append('The test Gone is not in this model, so it is left out.')
    path = rep.build_workbook(data, {'statuses': cr.OPEN, 'images': True, 'combined': True}).save(
        os.path.join(folder, 'report.xlsx'))
    with warnings.catch_warnings():
        warnings.simplefilter('error')
        wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Summary', 'Ducts vs Beams', 'Pipes vs Walls', 'Sprinklers vs Ceilings',
                             'MEP Ducts Trays L2', 'All clashes']
    summary = wb['Summary']
    values = [[c.value for c in row] for row in summary.iter_rows()]
    assert values[0][0] == 'Clash report: St Mary Hospital'
    kpi_row = [i for i, r in enumerate(values) if r[0] == 'Clashes'][0]
    assert values[kpi_row + 1][:2] == [600, 420]
    header = [i for i, r in enumerate(values) if r[0] == 'Test'][0]
    assert values[header][:6] == ['Test', 'Type', 'Tolerance', 'Last run', 'Test status', 'Clashes']
    total_row = [r for r in values if r[0] == 'Total'][0]
    assert total_row[5] == 600 and total_row[11] == 420
    pipes = values[header + 2]
    assert pipes[4] == 'Out of date'
    assert summary.cell(row=header + 3, column=5).fill.fgColor.rgb.endswith('FFF4D6')      # warn wash
    assert summary.cell(row=header + 1, column=7).fill.fgColor.rgb.endswith('F9D3D3')      # New header
    link = summary.cell(row=header + 2, column=14)
    assert link.value == 'Open' and link.hyperlink.location == "'Ducts vs Beams'!A1"
    assert any(r[0] == 'Open clashes by trade pair' for r in values)
    assert any(r[0] and 'Gone' in str(r[0]) for r in values)
    assert summary.freeze_panes is None

    sheet = wb['Ducts vs Beams']
    first = [c.value for c in sheet[4]]
    assert first[:5] == ['Image', 'Clash', 'Group', 'Status', 'Distance (mm)']
    assert sheet.freeze_panes == 'C5'
    listed = len(cr.filter_rows(data['tests'][0]['rows'], cr.OPEN))
    assert sheet.auto_filter.ref == 'A4:V%d' % (4 + listed)
    assert len(sheet._images) == 5
    assert type(sheet._images[0].anchor).__name__ == 'TwoCellAnchor'
    statuses = set(sheet.cell(row=r, column=4).value for r in range(5, 5 + listed))
    assert statuses <= set(cr.OPEN)
    status_cell = sheet.cell(row=5, column=4)
    assert status_cell.fill.fgColor.rgb.endswith({'New': 'F9D3D3', 'Active': 'FCE1C7',
                                                  'Reviewed': 'D6E6F8'}[status_cell.value])
    assert sheet.cell(row=5, column=5).number_format == '#,##0.0'
    assert sheet.cell(row=5, column=8).value == u'Duct <Supply> & Return ± Café'
    assert sheet.cell(row=5, column=21).value.startswith('<script>')       # kept as text, not markup

    everything = wb['All clashes']
    assert everything.max_row == 3 + 420 and everything['A3'].value == 'Test'
    assert not everything._images


def test_no_images_and_empty_listing():
    openpyxl = _openpyxl()
    folder = _compat.tmpdir('rep-clash-excel-empty')
    data = fx.fake_data()
    for test in data['tests']:
        test['rows'] = [r for r in test['rows'] if r['status'] != 'Approved'] if test['name'] != 'Pipes vs Walls' \
            else test['rows']
    path = rep.build_workbook(data, {'statuses': ['Approved'], 'images': False, 'combined': False}).save(
        os.path.join(folder, 'approved.xlsx'))
    wb = openpyxl.load_workbook(path)
    assert 'All clashes' not in wb.sheetnames
    empty = wb['Ducts vs Beams']
    assert empty['A4'].value == 'Clash' and empty['A5'].value == 'No clashes with these statuses.'
    full = wb['Pipes vs Walls']
    assert full['A4'].value == 'Clash' and full.freeze_panes == 'B5'


def test_identical_on_both_engines():
    """The same data gives the same sheets, strings and styles on CPython and IronPython (no
    dict-order or float-repr drift): the CPython run records a digest that the IronPython
    run of this test then has to match."""
    import io
    import zipfile
    wb = rep.build_workbook(fx.fake_data(), {'statuses': cr.OPEN, 'images': False})
    archive = zipfile.ZipFile(io.BytesIO(wb.to_bytes()))
    parts = sorted(n for n in archive.namelist() if n.startswith('xl/') and n.endswith('.xml'))
    fx.cross_engine('clash_excel', b''.join(archive.read(n) for n in parts))


def test_run_through_fake_navisworks():
    doc, _a, _b = fx.fake_clash_doc()
    with fx.Redirected('rep-clash-excel-run') as folder:
        with fx.nav.FakeNav(doc):
            result = reports.run('clash_excel', {'images': False, 'folder': folder}, doc=doc)
            missing = reports.run('clash_excel', {'images': False, 'folder': folder, 'tests': ['Nope']}, doc=doc)
            partly = reports.run('clash_excel', {'images': False, 'folder': folder,
                                                 'tests': ['Ducts vs Walls', 'Nope']}, doc=doc)
    assert result['level'] == 'success', result
    assert result['message'] == 'Saved the clash report'
    # the toast says where the file is: the full path, not just its name
    assert result['detail'] == '3 clashes from 2 tests, in %s' % result['paths'][0]
    assert os.path.isabs(result['paths'][0])
    assert os.path.isfile(result['paths'][0]) and result['paths'][0].endswith('.xlsx')
    assert missing['level'] == 'warning' and 'Nope' in missing['detail'] and not missing['paths']
    assert partly['level'] == 'warning' and partly['message'].endswith('some tests were not found')
