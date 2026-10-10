# -*- coding: utf-8 -*-
"""rep_clash_trends: series, burn-down text, workbook and CSV from constructed points, and
run() over snapshots recorded through the fake Navisworks layer."""
import datetime
import io
import os

import _compat

import test_clashreport as fx
from namo import clashrep_history as hist
from namo import clashreport as cr
from namo import reports
from namo import rep_clash_trends as rep


def _counts(**by_status):
    return {'total': sum(by_status.values()), 'by_status': by_status}


def _points():
    entries = []
    for day, (new, resolved) in enumerate([(40, 0), (34, 9), (30, 15), (21, 22)]):
        when = datetime.datetime(2026, 9, 1 + day * 7, 9, 0)
        counts = {'Ducts vs Beams': _counts(New=new, Resolved=resolved)}
        if day:
            counts['Pipes vs Walls'] = _counts(Active=10 - day)
        entries.append({'id': 'h%d' % day, 'time': when.strftime('%Y-%m-%dT%H:%M:%S'),
                        'label': 'Week %d' % (day + 1) if day == 2 else when.strftime('%Y-%m-%d %H:%M:%S'),
                        'counts': counts})
    current = {'time': datetime.datetime(2026, 10, 5, 9), 'counts': {'Ducts vs Beams': _counts(New=18, Resolved=25),
                                                                    'Pipes vs Walls': _counts(Active=5)}}
    return hist.trend_points(entries, None, current)


def test_contract():
    assert reports.title_of('clash_trends') == rep.TITLE
    config = rep.default_config()
    assert config['series'] == 'total' and config['current'] is True
    assert rep.describe(config) == 'Every test with results, open and resolved, with the current state'


def test_series_messages_and_rows():
    points, names = _points()
    assert names == ['Ducts vs Beams', 'Pipes vs Walls']
    total = rep.series_of(points, names, 'total')
    assert [n for n, _v in total] == ['Open', 'Resolved']
    assert total[0][1] == [40, 43, 38, 28, 23] and total[1][1] == [0, 9, 15, 22, 25]
    per = rep.series_of(points, names, 'per_test')
    assert per[1] == ('Pipes vs Walls', [None, 9, 8, 7, 5])
    assert rep.trend_message(points) == '23 open now, down 17 since 2026-09-01'
    assert rep.trend_message(points[:1]) == '40 open now'
    assert rep.trend_message([]) == 'No points to draw'
    labels = hist.point_labels(points)
    rows = rep.table_rows(points, labels, [None, (5, 2), (1, 6), (0, 10), (2, 7)])
    assert rows[1][-1] == '+3' and rows[2][-1] == '-5' and rows[0][-1] == ''
    assert rows[2][1] == 'Week 3' and rows[1][1] == '' and rows[-1][1] == 'Current state'
    assert rows[0][2] == '1 of 2'
    assert rep.span_text(points, names) == ('4 snapshots from 2026-09-01 to 2026-09-22 and the current state, '
                                            '2 tests.')
    head = rep.output_html_head(points, names)
    assert 'Clash trends' in head and 'went from 40 to 23' in head
    table = rep.output_html_table(points, labels, [None] * 5)
    assert 'Burn-down' in table and '1 of 2' in table and 'plots as zero' in table


def test_workbook_and_csv():
    points, names = _points()
    labels = hist.point_labels(points)
    flow = [None, (5, 2), (1, 6), (0, 10), (2, 7)]
    folder = _compat.tmpdir('rep-clash-trends')
    csv_path = rep.write_csv(os.path.join(folder, 'trend.csv'), points, labels, flow)
    with io.open(csv_path, encoding='utf-8-sig') as handle:
        lines = handle.read().splitlines()
    assert lines[0] == ','.join(rep.HEADERS)
    assert lines[1] == '2026-09-01,,1 of 2,40,0,40,,,' and lines[-1].startswith('Now,Current state,2,23,25,')
    xlsx_path = rep.build_workbook(points, names, labels, flow, 'Hospital').save(os.path.join(folder, 'trend.xlsx'))
    if _compat.IRONPYTHON:
        return
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    assert wb.sheetnames == ['Trend', 'Open by test']
    trend = wb['Trend']
    header = [c.value for c in trend[4]]
    assert header == rep.HEADERS
    assert [trend.cell(row=r, column=4).value for r in range(5, 10)] == [40, 43, 38, 28, 23]
    per = wb['Open by test']
    assert per['A3'].value == 'When' and per['B4'].value == 40 and per['C4'].value is None and per['C5'].value == 9


def test_run_over_recorded_snapshots():
    doc, test_a, test_b = fx.fake_clash_doc()
    with fx.Redirected('rep-clash-trends-run') as folder:
        doc.FileName = os.path.join(folder, 'Hospital.nwf')
        with fx.nav.FakeNav(doc):
            nothing = reports.run('clash_trends', {}, doc=doc, interactive=False)
            cr.record_snapshot([test_a, test_b], 'First', doc)
            test_a.Children[0].Status = fx.nav._Enum('Resolved')
            cr.record_snapshot([test_a, test_b], 'Second', doc)
            test_b.Children[0].Status = fx.nav._Enum('Approved')
            result = reports.run('clash_trends', {'folder': folder, 'csv': True}, doc=doc, interactive=False)
            only = reports.run('clash_trends', {'folder': folder, 'current': False, 'xlsx': False, 'csv': True,
                                                'tests': ['Ducts vs Walls'], 'name_template': 'only'},
                               doc=doc, interactive=False)
    assert nothing['message'] == 'No clash snapshots yet'
    assert result['level'] == 'success', result
    assert result['message'].startswith('1 open now, down 2 since ')
    assert [os.path.splitext(p)[1] for p in result['paths']] == ['.xlsx', '.csv']
    with io.open(result['paths'][1], encoding='utf-8-sig') as handle:
        rows = handle.read().splitlines()
    assert len(rows) == 4 and rows[2].split(',')[6:8] == ['0', '1'] and rows[3].split(',')[6:8] == ['0', '1']
    assert only['paths'][0].endswith('only.csv')
