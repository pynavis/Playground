# -*- coding: utf-8 -*-
"""rep_coordination: the namo.reports contract, row selection, the HTML and the workbook
(structure checked with openpyxl under CPython), and a full non-interactive run against a
sample register."""
import datetime
import io
import os

import _compat

from namo import issues, issuestore, paths, reports, rep_coordination

NOW = datetime.datetime.now().replace(microsecond=0)
DOC = r'C:\Projects\Tower\Coordination.nwf'


class FakeDoc(object):
    def __init__(self, path):
        self.FileName = path


def _seed():
    s = issuestore.IssueStore(DOC)
    day = datetime.timedelta(days=1)
    s.create({'title': 'Duct hits beam', 'trade': 'Mechanical', 'priority': 'Critical',
              'due': (NOW - 3 * day).strftime('%Y-%m-%d')}, prefix='RFI-', user='amol')
    s.create({'title': 'Sprinkler clearance', 'trade': 'Fire Protection',
              'due': (NOW + 2 * day).strftime('%Y-%m-%d')}, prefix='VER-', user='amol')
    s.create({'title': 'Cable tray agreed', 'trade': 'Electrical', 'status': 'Closed'}, user='amol')
    s.create({'title': 'Not an issue', 'status': 'Void'}, user='amol')
    return s


def test_contract():
    config = rep_coordination.default_config()
    assert config['folder'] == '' and config['name_template']
    assert rep_coordination.TITLE == reports.title_of('coordination')
    assert rep_coordination.describe(config) == 'HTML and XLSX, every trade, thumbnails'
    assert rep_coordination.describe(dict(config, trades=['A'], thumbnails=False,
                                          formats=['html'])) == 'HTML, 1 trade'


def test_select_rows_and_headline():
    rows = [issues.normalise({'id': 'A-1', 'title': 'x', 'trade': 'Mechanical'}),
            issues.normalise({'id': 'A-2', 'title': 'y', 'status': 'Void'})]
    assert [r['id'] for r in rep_coordination.select_rows(rows, {})] == ['A-1']
    assert len(rep_coordination.select_rows(rows, {'include_void': True})) == 2
    assert rep_coordination.select_rows(rows, {'trades': ['Electrical']}) == []
    summary = issues.summary(rows)
    assert rep_coordination.headline(summary) == 'Coordination report: 1 open, 0 overdue'


def test_run_writes_html_and_workbook():
    folder = _compat.tmpdir('g8-rep-coordination')
    previous = paths.redirect(folder)
    try:
        _seed()
        out = os.path.join(folder, 'out')
        result = reports.run('coordination', {'folder': out, 'thumbnails': True,
                                              'name_template': '{report} test'},
                             doc=FakeDoc(DOC), interactive=False)
        assert result['level'] == 'success', result
        assert result['message'] == 'Coordination report: 2 open, 1 overdue'
        assert [os.path.basename(p) for p in result['paths']] == ['Coordination Report test.html',
                                                                 'Coordination Report test.xlsx']
        with io.open(result['paths'][0], encoding='utf-8') as handle:
            page = handle.read()
        for text in ('Overdue', 'Due this week', 'Mechanical', 'Recently closed', 'RFI-001'):
            assert text in page, text
        assert 'Not an issue' not in page
        unsaved = reports.run('coordination', {}, doc=FakeDoc(''), interactive=False)
        assert unsaved['level'] == 'info' and unsaved['message'] == 'Save the model first'
        if _compat.IRONPYTHON:
            return
        import openpyxl
        wb = openpyxl.load_workbook(result['paths'][1])
        assert wb.sheetnames == ['Summary', 'Open issues', 'Overdue', 'Recently closed', 'All issues']
        all_rows = list(wb['All issues'].iter_rows(values_only=True))
        assert all_rows[0][0] == 'Issue' and len(all_rows) == 4
        overdue = list(wb['Overdue'].iter_rows(values_only=True))
        assert overdue[1][0] == 'RFI-001' and overdue[1][8] == 3
    finally:
        paths.redirect(previous)


def test_empty_register():
    folder = _compat.tmpdir('g8-rep-coordination-empty')
    previous = paths.redirect(folder)
    try:
        result = reports.run('coordination', {}, doc=FakeDoc(DOC), interactive=False)
        assert result['level'] == 'info' and result['message'] == 'No issues to report'
    finally:
        paths.redirect(previous)
