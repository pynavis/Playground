# -*- coding: utf-8 -*-
"""namo.health and namo.rep_health_check: scoring, the outlier search, checks, writers,
and the collection pass through the fake Navisworks layer."""
import io
import json
import os

import _compat

from namo import health


# ---- settings and maths ----------------------------------------------------------------


def test_settings_coerce_and_order():
    cfg = health.settings({'warn_km': '5', 'fail_km': 2, 'outlier_factor': 1, 'age_days': '0',
                           'item_budget': 'x', 'formats': ['xlsx', 'pdf', 'html'], 'folder': None})
    assert cfg['warn_km'] == 5.0 and cfg['fail_km'] == 5.0
    assert cfg['outlier_factor'] == 2.0 and cfg['age_days'] == 1
    assert cfg['item_budget'] == health.DEFAULTS['item_budget']
    assert cfg['formats'] == ['html', 'xlsx'] and cfg['folder'] == ''
    assert health.settings()['formats'] is not health.DEFAULTS['formats']


def test_median_and_points():
    assert health.median([]) is None
    assert health.median([3, 1, 2]) == 2.0
    assert health.median([4, 1, 2, 3]) == 2.5
    assert health.median_point([(0, 0, 0), (10, 1, 5), (2, 100, 4), None]) == (2.0, 1.0, 4.0)
    assert health.median_point([]) is None
    box = ((-1, -2, -3), (1, 2, 3))
    assert _compat.approx(health.farthest_corner(box, (0, 0, 0)), (1 + 4 + 9) ** 0.5)
    assert _compat.approx(health.farthest_corner(box, (1, 2, 3)), (4 + 16 + 36) ** 0.5)
    assert health.box_center(box) == (0.0, 0.0, 0.0)
    assert health.box_center(None) is None and health.box_diagonal(None) is None


def test_stride_sampler_keeps_an_even_sample():
    sampler = health.StrideSampler(100)
    for value in range(10000):
        if sampler.offer():
            sampler.add(value)
    assert 100 <= len(sampler.kept) < 200
    stride = sampler.stride
    assert all(v % stride == 0 for v in sampler.kept)
    assert sampler.kept[0] == 0 and sampler.kept == sorted(sampler.kept)
    small = health.StrideSampler(100)
    for value in range(50):
        if small.offer():
            small.add(value)
    assert small.kept == list(range(50))


def test_spread():
    assert health.spread(range(10), 20) == list(range(10))
    assert health.spread(range(100), 4) == [0, 25, 50, 75]
    assert health.spread([1, 2], 0) == []


class Node(object):
    def __init__(self, name, box=None, children=()):
        self.name = name
        self.box = box
        self.Children = list(children)


def _cube(cx, size=1.0):
    return ((cx - size, -size, -size), (cx + size, size, size))


def _union(nodes):
    boxes = [n.box for n in nodes if n.box is not None]
    return (tuple(min(b[0][i] for b in boxes) for i in range(3)),
            tuple(max(b[1][i] for b in boxes) for i in range(3)))


def _tree(far_at=1000.0):
    near = [Node('n%d' % i, _cube(i * 2.0)) for i in range(10)]
    near_group = Node('near', _union(near), near)
    far = Node('far', _cube(far_at))
    deep = Node('deep', _union([far]), [far])
    mixed = Node('mixed', _union([deep, Node('x', _cube(3.0))]), [Node('x2', _cube(3.0)), deep])
    root = Node('root', _union([near_group, mixed]), [near_group, mixed])
    return root


def test_find_far_skips_subtrees_inside_the_radius():
    root = _tree()
    calls = []

    def box_of(node):
        calls.append(node.name)
        return node.box
    hits, sampled = health.find_far([root], lambda n: n.Children, box_of, lambda n: not n.Children,
                                    (5.0, 0.0, 0.0), 100.0)
    assert [h[0].name for h in hits] == ['far'] and not sampled
    assert _compat.approx(hits[0][1], 995.0)
    assert 'n3' not in calls and 'near' in calls          # the near group was skipped whole
    clean, sampled = health.find_far([_tree(far_at=50.0)], lambda n: n.Children, lambda n: n.box,
                                     lambda n: not n.Children, (5.0, 0.0, 0.0), 100.0)
    assert clean == [] and not sampled


def test_find_far_budget_and_limit_are_honest():
    leaves = [Node('f%d' % i, _cube(1000.0 + i)) for i in range(5)]
    root = Node('root', _union(leaves), leaves)
    hits, sampled = health.find_far([root], lambda n: n.Children, lambda n: n.box,
                                    lambda n: not n.Children, (0, 0, 0), 10.0, limit=3)
    assert len(hits) == 3 and sampled
    hits, sampled = health.find_far([root], lambda n: n.Children, lambda n: n.box,
                                    lambda n: not n.Children, (0, 0, 0), 10.0, node_budget=2)
    assert len(hits) == 1 and sampled
    hits, sampled = health.find_far([root], lambda n: n.Children, lambda n: n.box,
                                    lambda n: not n.Children, (0, 0, 0), 10.0, limit=5)
    assert len(hits) == 5 and not sampled
    assert [h[0].name for h in hits] == ['f4', 'f3', 'f2', 'f1', 'f0']      # farthest first
    hits, _s = health.find_far([Node('nobox')], lambda n: [], lambda n: None, lambda n: True,
                               (0, 0, 0), 1.0)
    assert hits == []


def test_formatting():
    assert health.fmt_m(412345) == '412 km'
    assert health.fmt_m(12400) == '12.4 km'
    assert health.fmt_m(84.2) == '84 m'
    assert health.fmt_m(3.25) == '3.2 m' or health.fmt_m(3.25) == '3.3 m'
    assert health.fmt_m(0.35) == '35 cm'
    assert health.fmt_m(0.004) == '4 mm'
    assert health.fmt_m(None) == ''
    assert health.num_text(10.0) == '10' and health.num_text(2.5) == '2.5'
    assert health.age_days_text(10) == 'today'
    assert health.age_days_text(86400 * 1.5) == '1 day ago'
    assert health.age_days_text(86400 * 1200) == '1,200 days ago'


# ---- checks ---------------------------------------------------------------------------------


def _facts(**overrides):
    facts = {'index': 0, 'name': 'ARCH.nwc', 'kind': 'RVT', 'units': 'Meters', 'doc_units': 'Meters',
             'diag_m': 80.0, 'center_m': 50.0, 'items': 1000, 'geometry': 600, 'walk_sampled': False,
             'd_med_m': 10.0, 'outliers': [], 'outliers_sampled': False, 'outliers_searched': True,
             'identity': {'checked': 300, 'with_ids': 300}, 'duplicate_of': [],
             'transformed': False, 'now': 1000000.0,
             'stale': {'status': 'Current', 'level': 'pass', 'detail': 'ok', 'cache_time': 1000000.0 - 3600},
             'error': None}
    facts.update(overrides)
    return facts


def _levels(checks):
    return dict((c['id'], c['level']) for c in checks)


def test_clean_model_scores_100():
    model = health.scored(_facts())
    assert model['score'] == 100 and model['grade'] == 'Good'
    assert model['fails'] == 0 and model['warns'] == 0
    levels = _levels(model['checks'])
    assert levels['units'] == 'pass' and levels['origin'] == 'pass' and levels['items'] == 'info'
    assert [c['id'] for c in model['checks']] == ['units', 'size', 'origin', 'outliers', 'geometry',
                                                  'identity', 'duplicate', 'transform', 'age', 'items']


def test_origin_size_and_units_levels():
    assert _levels(health.evaluate(_facts(center_m=12000.0)))['origin'] == 'warn'
    assert _levels(health.evaluate(_facts(center_m=150000.0)))['origin'] == 'fail'
    custom = {'warn_km': 1, 'fail_km': 2}
    assert _levels(health.evaluate(_facts(center_m=1500.0), custom))['origin'] == 'warn'
    assert _levels(health.evaluate(_facts(diag_m=6000.0)))['size'] == 'warn'
    assert _levels(health.evaluate(_facts(diag_m=60000.0)))['size'] == 'fail'
    assert _levels(health.evaluate(_facts(diag_m=0.01)))['size'] == 'warn'
    units = health.evaluate(_facts(units='Millimeters'))
    assert _levels(units)['units'] == 'info' and 'Millimeters' in units[0]['detail']


def test_outlier_identity_duplicate_transform_age_levels():
    far = [{'name': 'Stray', 'distance_m': 412000.0, 'path': 'a > b'}]
    checks = health.evaluate(_facts(outliers=far, outliers_sampled=True))
    out = [c for c in checks if c['id'] == 'outliers'][0]
    assert out['level'] == 'warn' and '412 km' in out['detail'] and 'may be more' in out['detail']
    assert _levels(health.evaluate(_facts(outliers_searched=False)))['outliers'] == 'info'
    assert _levels(health.evaluate(_facts(identity={'checked': 300, 'with_ids': 200})))['identity'] == 'info'
    assert _levels(health.evaluate(_facts(identity={'checked': 300, 'with_ids': 10})))['identity'] == 'warn'
    dwg = health.evaluate(_facts(kind='DWG', identity={'checked': 300, 'with_ids': 0}))
    assert _levels(dwg)['identity'] == 'info' and 'DWG' in [c for c in dwg if c['id'] == 'identity'][0]['detail']
    assert _levels(health.evaluate(_facts(kind='RVT', identity={'checked': 300, 'with_ids': 0})))['identity'] == 'warn'
    dup = health.evaluate(_facts(index=2, duplicate_of=[0, 2]))
    assert _levels(dup)['duplicate'] == 'fail'
    assert 'model 1' in [c for c in dup if c['id'] == 'duplicate'][0]['detail']
    assert _levels(health.evaluate(_facts(transformed=True)))['transform'] == 'warn'
    old = _facts(stale={'status': 'Current', 'detail': 'ok', 'cache_time': 1000000.0 - 86400 * 45})
    assert _levels(health.evaluate(old))['age'] == 'info'
    late = _facts(stale={'status': 'Out of date', 'detail': 'x.rvt was saved later', 'cache_time': 1.0})
    assert _levels(health.evaluate(late))['age'] == 'warn'
    lost = _facts(stale={'status': 'File not found', 'detail': 'gone', 'cache_time': None})
    assert _levels(health.evaluate(lost))['age'] == 'warn'
    assert _levels(health.evaluate(_facts(stale=None)))['age'] == 'info'


def test_empty_model_and_read_error():
    assert _levels(health.evaluate(_facts(geometry=0)))['geometry'] == 'fail'
    assert _levels(health.evaluate(_facts(geometry=0, walk_sampled=True)))['geometry'] == 'warn'
    errored = health.scored(_facts(error='boom'))
    assert len(errored['checks']) == 1 and errored['checks'][0]['level'] == 'fail'
    assert errored['score'] == 75
    counted = health.evaluate(_facts(walk_sampled=True, items=400000))
    assert 'At least 400,000' in [c for c in counted if c['id'] == 'items'][0]['detail']


def test_score_and_grade():
    def checks(fails, warns):
        return ([{'level': 'fail'}] * fails) + ([{'level': 'warn'}] * warns) + [{'level': 'pass'}]
    assert health.score(checks(0, 0)) == 100
    assert health.score(checks(1, 1)) == 67
    assert health.score(checks(5, 0)) == 0
    assert health.grade(85) == 'Good' and health.grade(84) == 'Fair'
    assert health.grade(60) == 'Fair' and health.grade(59) == 'Poor'
    assert health.grade_level('Poor') == 'fail'


def test_summary_messages():
    good = health.scored(_facts())
    fair = health.scored(_facts(index=1, name='MEP.nwc', center_m=20000.0, transformed=True, units='Feet'))
    poor = health.scored(_facts(index=2, name='STR.nwc', center_m=500000.0, diag_m=9000.0, geometry=0))
    assert health.summary([]) [0] == 'info'
    level, message, detail = health.summary([good])
    assert level == 'success' and message == 'The model looks healthy' and detail == 'Score 100.'
    level, message, _d = health.summary([good, health.scored(_facts(index=1, name='B'))])
    assert level == 'success' and message == 'All 2 models look healthy'
    level, message, detail = health.summary([good, fair])
    assert level == 'warning' and message == 'Health check: 1 model needs attention'
    assert detail == 'MEP.nwc 84 (Fair).'
    level, message, detail = health.summary([good, fair, poor])
    assert level == 'warning' and '2 models need' in message
    assert detail == 'STR.nwc 50 (Poor), MEP.nwc 84 (Fair).'
    view = health.overview([good, fair, poor])
    assert view['models'] == 3 and view['fails'] == 2 and view['average'] == int(round((100 + 84 + 50) / 3.0))


def test_json_safe_and_no_dashes():
    models = [health.scored(_facts(outliers=[{'name': 'S', 'distance_m': 1e6, 'path': 'x'}]))]
    text = json.dumps(health.json_safe(models), sort_keys=True)
    assert 'Stray' not in text and '"score": ' in text
    for model in models:
        for check in model['checks']:
            assert u'\u2014' not in check['detail'] and u'\u2013' not in check['detail']
            assert check['detail'][:1].isupper() or check['detail'][:1].isdigit()


# ---- writers and output html ------------------------------------------------------------------------


def _models():
    far = [{'name': '<Stray & co>', 'distance_m': 412000.0, 'path': 'ARCH > L1 > Stray'}]
    return [health.scored(_facts(outliers=far, file=r'C:\M\ARCH.nwc')),
            health.scored(_facts(index=1, name='MEP & co.nwc', center_m=20000.0))]


def test_write_html_escapes_and_lists_models():
    folder = _compat.tmpdir('health_html' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    path = health.write_html(os.path.join(folder, 'h.html'), _models(), subtitle='Hospital',
                             meta=[('Document', 'Hospital.nwf')])
    with io.open(path, encoding='utf-8') as handle:
        text = handle.read()
    assert 'MEP &amp; co.nwc' in text and '&lt;Stray &amp; co&gt;' in text
    assert '<Stray' not in text
    assert 'Far-away objects' in text and 'Distance from origin' in text


def test_write_xlsx_round_trip():
    folder = _compat.tmpdir('health_xlsx' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    path = health.write_xlsx(os.path.join(folder, 'h.xlsx'), _models(), author='me')
    assert os.path.getsize(path) > 0
    _compat.require_cpython('openpyxl round trip')
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Summary', 'Checks', 'Outliers']
    summary = [[c.value for c in row] for row in wb['Summary'].iter_rows()]
    assert summary[1][:3] == ['Model', 'Score', 'Grade']
    assert summary[2][:2] == ['ARCH.nwc', 92] and summary[3][:2] == ['MEP & co.nwc', 92]
    checks = list(wb['Checks'].iter_rows(values_only=True))
    assert len(checks) == 1 + 2 * 10
    outliers = list(wb['Outliers'].iter_rows(values_only=True))
    assert outliers[1][1] == '<Stray & co>' and outliers[1][2] == 412000.0


def test_card_html_links_and_escapes():
    model = _models()[0]
    html = health.card_html(model, link=lambda i: '<a class="pynavis-el">L%d</a>' % i)
    assert '<a class="pynavis-el">L0</a>' in html and 'Pass' in html
    plain = health.card_html(_models()[1])
    assert 'MEP &amp; co.nwc' in plain and 'Far-away objects</h3>' not in plain
    unlinked = health.card_html(model, link=lambda i: None)
    assert '&lt;Stray &amp; co&gt;' in unlinked
    assert 'Average score' in health.overview_html(_models())


# ---- the report module --------------------------------------------------------------------------


def test_report_config_is_json_safe_and_described():
    from namo import rep_health_check, reports
    config = rep_health_check.default_config()
    assert config['folder'] == '' and config['name_template'] == rep_health_check.NAME_TEMPLATE
    assert config['formats'] == ['html', 'xlsx'] and config['formats'] is not health.DEFAULTS['formats']
    json.dumps(config, sort_keys=True)
    text = rep_health_check.describe(config)
    assert 'HTML and Excel' in text and '10 km' in text
    assert rep_health_check.TITLE == reports.title_of('health_check')
    assert 'no file' in rep_health_check.describe(dict(config, formats=[]))


class _Ns(object):
    def __init__(self, **kw):
        self.__dict__.update(kw)


def test_report_with_no_models_says_what_to_do():
    from namo import reports
    empty = _Ns(Models=_Ns(Count=0), Units='Meters')
    result = reports.run('health_check', {}, doc=empty, interactive=False)
    assert result['level'] == 'info' and result['models'] == []
    assert 'Append models' in result['detail'] and result['paths'] == []


# ---- the API half through the fake layer (see test_items) ------------------------------------------

import test_items as nav  # noqa: E402


def _scattered_doc():
    """One model: 12 objects in a 20 m cluster, one stray object 500 km away."""
    objects = []
    for i in range(12):
        geo = nav.ModelItem('Part %d' % i, geometry=True,
                            box=((i * 1.5, 0, 0), (i * 1.5 + 1, 1, 3)))
        element = nav.ModelItem('Element %d' % i, [geo], props={('Element ID', 'Value'): 100 + i})
        geo.object = element
        objects.append(element)
    stray_geo = nav.ModelItem('Stray solid', geometry=True, box=((500000, 0, 0), (500001, 1, 1)))
    stray = nav.ModelItem('Stray', [stray_geo])
    stray_geo.object = stray
    level = nav.ModelItem('Level 1', objects + [stray])
    root = nav.ModelItem('SITE.nwc', [level])
    model = nav.Model(r'C:\NoSuchFolder\SITE.rvt', root)
    root.Model, root.HasModel = model, True
    empty_root = nav.ModelItem('EMPTY.nwc', [nav.ModelItem('Nothing')])
    empty = nav.Model(r'C:\NoSuchFolder\EMPTY.dwg', empty_root)
    empty_root.Model, empty_root.HasModel = empty, True
    return nav.Document([model, empty])


def test_collect_finds_the_stray_and_scores():
    doc = _scattered_doc()
    with nav.FakeNav(doc):
        reports = []
        models, outliers = health.collect(doc, {'outlier_factor': 20},
                                          progress=lambda f, label: reports.append((f, label)))
    assert [m['name'] for m in models] == ['SITE.rvt', 'EMPTY.dwg']
    site, empty = models
    assert site['items'] == 28 and site['geometry'] == 13 and not site['walk_sampled']
    assert [o['name'] for o in site['outliers']] == ['Stray']
    assert outliers[0][0].DisplayName == 'Stray'
    assert site['outliers'][0]['distance_m'] > 400000
    levels = _levels(site['checks'])
    assert levels['outliers'] == 'warn' and levels['origin'] == 'fail'
    assert site['identity'] == {'checked': 13, 'with_ids': 12} and levels['identity'] == 'pass'
    assert levels['age'] == 'warn'                  # the loaded file is not on this computer
    assert _levels(empty['checks'])['geometry'] == 'fail'
    assert reports and reports[-1] == (1.0, 'Done')
    json.dumps(health.json_safe(models), sort_keys=True)


def test_collect_budget_and_cancel():
    doc = _scattered_doc()
    with nav.FakeNav(doc):
        models, _outliers = health.collect(doc, {'item_budget': 1000})
        assert models[0]['walk_sampled'] is False

        class Cancelled(Exception):
            pass

        def stop(fraction, label):
            raise Cancelled()
        _compat.raises(Cancelled, health.collect, doc, None, stop)


def test_report_run_writes_both_files_through_the_fake_layer():
    from namo import paths, reports
    folder = _compat.tmpdir('health_report' + ('_ipy' if _compat.IRONPYTHON else '_cpy'))
    previous = paths.redirect(folder)
    try:
        doc = _scattered_doc()
        with nav.FakeNav(doc) as fake:
            import sys
            sys.modules['pynavis.doc'].get_filename = lambda d=None: ''
            sys.modules['pynavis.doc'].get_title = lambda d=None: 'Site'
            result = reports.run('health_check', {'folder': folder}, doc=doc, interactive=True)
    finally:
        paths.redirect(previous)
    assert result['level'] in ('warning', 'success'), result
    assert len(result['paths']) == 2 and all(os.path.exists(p) for p in result['paths'])
    assert os.path.basename(result['paths'][0]).startswith('Untitled - Health check - ')
    assert result['models'][0]['name'] == 'SITE.rvt'
    assert result['outlier_items'][0][0].DisplayName == 'Stray'
    quiet = None
    with nav.FakeNav(_scattered_doc()) as fake:
        import sys
        sys.modules['pynavis.doc'].get_filename = lambda d=None: ''
        sys.modules['pynavis.doc'].get_title = lambda d=None: 'Site'
        quiet = reports.run('health_check', {'folder': folder, 'formats': []},
                            doc=fake.doc, interactive=False)
    assert quiet['paths'] == [] and 'outlier_items' not in quiet


# ---- refit: far from the origin is judged by how true the geometry still reads -----------


def test_origin_far_out_but_precise_is_not_a_failure():
    far_m = 1370000.0 * 0.3048                  # the field model: 418 km out, in feet
    plain = health.evaluate(_facts(center_m=far_m))
    assert _levels(plain)['origin'] == 'fail'   # distance alone, when nothing was measured
    fine = health.evaluate(_facts(center_m=far_m, noise_m=1e-6 * 0.3048))
    origin = [c for c in fine if c['id'] == 'origin'][0]
    assert origin['level'] == 'info' and 'shared coordinates' in origin['detail']
    assert 'under 0.001 mm' in origin['detail'] and '418 km' in origin['detail']
    assert health.scored(_facts(center_m=far_m, noise_m=1e-7))['score'] == 100
    coarse = [c for c in health.evaluate(_facts(center_m=far_m, noise_m=0.004)) if c['id'] == 'origin'][0]
    assert coarse['level'] == 'warn' and '4 mm' in coarse['detail']
    bad = [c for c in health.evaluate(_facts(center_m=far_m, noise_m=0.05)) if c['id'] == 'origin'][0]
    assert bad['level'] == 'fail' and '50 mm' in bad['detail']
    # near the origin the noise is not even looked at
    assert _levels(health.evaluate(_facts(center_m=50.0, noise_m=0.05)))['origin'] == 'pass'
    assert health.fmt_fine(0.0000004) == 'under 0.001 mm' and health.fmt_fine(0.0004) == '0.4 mm'
    assert health.fmt_fine(0.25) == '25 cm'


def test_measured_noise_takes_the_worst_readable_object():
    from namo import mesh
    readings = {'a': {'tris': [1], 'noise': 0.001}, 'b': {'tris': [], 'noise': 9.0},
                'c': {'tris': [1], 'noise': 0.003}}
    original = mesh.item_triangles
    mesh.item_triangles = lambda obj, budget=0: readings[obj]
    try:
        assert health.measured_noise(['a', 'b', 'c']) == 0.003
        assert health.measured_noise(['b']) is None
        assert health.measured_noise([]) is None
    finally:
        mesh.item_triangles = original
