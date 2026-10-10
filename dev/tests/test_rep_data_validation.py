# -*- coding: utf-8 -*-
"""rep_data_validation: contract, rule file choice, messages and a non-interactive run
(starter rules written into a redirected data root) through the fake layer."""
import io
import json
import os

import _compat
import test_propindex as pf

from namo import paths, reports, validator, rep_data_validation as rep


def test_contract_and_helpers():
    config = rep.default_config()
    assert 'folder' in config and 'name_template' in config
    assert rep.TITLE == reports.title_of('data_validation')
    assert rep.describe(config) == 'every rule file, whole model'
    sets = [{'file': 'A.json'}, {'file': 'b.json'}]
    assert rep.choose_rulesets(sets, []) == sets
    assert rep.choose_rulesets(sets, ['B.JSON']) == [sets[1]]


def test_messages():
    def result(checked, failed):
        return {'rule': 'r', 'checked': checked, 'passed': checked - failed, 'failed': failed}
    assert rep.message_for([result(0, 0)], 3, 'objects') == ('info', 'No rule applied to the 3 objects')
    assert rep.message_for([result(2, 0), result(1, 0)], 3, 'objects') == ('success',
                                                                          'All 2 rules passed on 3 objects')
    assert rep.message_for([result(2, 1), result(1, 0)], 1, 'geometry') == (
        'warning', '1 of 2 rules failed: 1 failure')


def test_run_with_rule_folder():
    root = _compat.tmpdir('g2_rep_validation')
    previous = paths.redirect(root)
    try:
        folder = validator.rules_folder()
        with io.open(os.path.join(folder, 'mine.json'), 'w', encoding='utf-8') as handle:
            handle.write(json.dumps({'name': 'Mine', 'rules': [
                {'name': 'Walls have a mark', 'applies_to': ['Element > Category = Walls'],
                 'require': ['Element > Mark is not empty']}]}))
        doc, n = pf.tree()
        with pf.Fake(doc):
            out = os.path.join(root, 'v.xlsx')
            result = rep.run(dict(rep.default_config(), rule_files=['mine.json'], path=out), doc=doc,
                             interactive=False)
            assert result['level'] == 'warning', result
            assert result['message'] == '1 of 1 rule failed: 1 failure'
            assert result['paths'] == [out] and os.path.exists(out)
            assert result['totals']['checked'] == 2
            assert not os.path.exists(os.path.join(folder, validator.STARTER_FILE))
            nothing = rep.run(dict(rep.default_config(), rule_files=['none.json'], path=out), doc=doc,
                              interactive=False)
            assert nothing['message'] == 'No rules to run'
        try:
            import openpyxl
        except ImportError:
            return
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ['Summary', 'Failures']
        failures = wb['Failures']
        assert failures.cell(2, 1).value == 'Walls have a mark' and failures.cell(2, 3).value == 'Wall 2'
        assert failures.cell(2, 7).value == 'Mark is empty'
    finally:
        paths.redirect(previous)


def test_starter_written_into_an_empty_folder():
    root = _compat.tmpdir('g2_rep_validation_starter')
    previous = paths.redirect(root)
    try:
        doc, n = pf.tree()
        with pf.Fake(doc):
            result = rep.run(dict(rep.default_config(), path=os.path.join(root, 's.xlsx')), doc=doc,
                             interactive=False)
            folder = result['folder']
            assert os.path.exists(os.path.join(folder, validator.STARTER_FILE))
            assert result['totals']['rules'] == len(validator.STARTER['rules'])
    finally:
        paths.redirect(previous)


def test_interactive_run_behind_progress_leaves_its_page_for_later():
    """Nothing prints while a progress window is up: the page, with its element links
    still unbuilt, comes back in result['output']."""
    root = _compat.tmpdir('g2_rep_validation_page')
    previous = paths.redirect(root)
    try:
        folder = validator.rules_folder()
        with io.open(os.path.join(folder, 'mine.json'), 'w', encoding='utf-8') as handle:
            handle.write(json.dumps({'name': 'Mine', 'rules': [
                {'name': 'Walls have a mark', 'applies_to': ['Element > Category = Walls'],
                 'require': ['Element > Mark is not empty']}]}))
        doc, n = pf.tree()
        with pf.Fake(doc):
            result = rep.run(dict(rep.default_config(), rule_files=['mine.json']), doc=doc,
                             interactive=True, progress=lambda fraction, label: True)
        assert result['message'] == '1 of 1 rule failed: 1 failure', result
        page = result['output']
        kinds = [c[0] for c in page.calls]
        assert kinds == ['html', 'html', 'later']          # card, summary, failure links at show time
    finally:
        paths.redirect(previous)
