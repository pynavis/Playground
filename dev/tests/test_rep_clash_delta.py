# -*- coding: utf-8 -*-
"""rep_clash_delta: output HTML and workbook from a constructed delta, and run() against a
snapshot recorded through the fake Navisworks layer after the clashes changed."""
import datetime
import os

import _compat

import test_clashreport as fx
from namo import clashrep_history as hist
from namo import clashreport as cr
from namo import reports
from namo import rep_clash_delta as rep


def _delta():
    data = fx.fake_data()
    new_by_test = dict((t['name'], t['rows']) for t in data['tests'])
    old = {}
    for test in data['tests']:
        rows = []
        for row in test['rows']:
            if row['id'] % 7 == 0:
                continue                                  # appeared since: new
            copy = dict(row)
            if row['id'] % 5 == 0 and row['status'] == 'Resolved':
                copy['status'] = 'Active'                 # was open: fixed by status
            rows.append(copy)
        rows.append(dict(test['rows'][1], key='gone|x@1,1,1', status='New', name='Old clash'))
        old[test['name']] = rows
    del old['MEP: Ducts / Trays [L2]']
    d = rep.sort_new(hist.delta(old, new_by_test, [t['name'] for t in data['tests']]))
    return data, d


def test_contract():
    assert reports.title_of('clash_delta') == rep.TITLE
    config = rep.default_config()
    assert config['baseline'] == 'auto' and config['folder'] == ''
    assert rep.describe(config) == 'Every test with results against the latest snapshot before today'
    assert rep.describe({'tests': ['A'], 'baseline': 'Meeting 4'}) == 'A against Meeting 4'


def test_delta_summary_rows_and_html():
    data, d = _delta()
    rows = rep.summary_rows(d)
    assert rows[-1][0] == 'Total' and len(rows) == 5
    assert rows[3][1] is None and rows[3][3] is None                  # MEP test had no baseline
    assert rows[-1][3] == d['totals']['new'] and rows[-1][4] == d['totals']['fixed']
    first = d['tests'][0]['new']
    assert [r['dist'] for r in first] == sorted(r['dist'] for r in first)
    entry = {'id': 'h1', 'label': 'Meeting 4', 'time': '2026-10-04T17:30:00', 'counts': {}}
    assert rep.snapshot_line(entry, data['generated']) == 'Snapshot Meeting 4 (2026-10-04 17:30), 1 day ago.'
    html = rep.output_html(d, entry, data, link=lambda row: '<a class="pynavis-el">Show</a>')
    assert 'Clash delta: now against the snapshot' in html
    assert '<a class="pynavis-el">Show</a>' in html
    assert 'Not in the snapshot' in html and 'MEP: Ducts / Trays [L2]' in html
    assert 'Old clash' in html and 'Gone' in html
    assert '<script>' not in html
    assert rep.change_text(-3) == 'down 3' and rep.change_text(0) == 'no change' and rep.change_text(None) == ''
    quiet = hist.delta({'A': []}, {'A': []}, ['A'])
    assert 'Nothing changed since the snapshot.' in rep.output_html(quiet, entry, data)


def test_delta_workbook():
    _compat.require_cpython('openpyxl reads the workbook under CPython')
    import openpyxl
    data, d = _delta()
    entry = {'id': 'h1', 'label': '', 'time': '2026-10-04T17:30:00', 'counts': {}}
    folder = _compat.tmpdir('rep-clash-delta-wb')
    path = rep.build_workbook(d, entry, data).save(os.path.join(folder, 'delta.xlsx'))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['Summary', 'New', 'Fixed', 'Status changed', 'Still open']
    new = wb['New']
    assert new.max_row == 3 + d['totals']['new'] and not d['tests'][3]['new']
    assert new.auto_filter.ref.startswith('A3:M')
    fixed = wb['Fixed']
    hows = set(fixed.cell(row=r, column=4).value for r in range(4, fixed.max_row + 1))
    assert hows <= set(['Gone', 'Resolved', 'Approved']) and 'Gone' in hows
    assert wb['Summary']['A4'].value == 'Test'


def test_run_against_a_recorded_snapshot():
    doc, test_a, test_b = fx.fake_clash_doc()
    with fx.Redirected('rep-clash-delta-run') as folder:
        doc.FileName = os.path.join(folder, 'Hospital.nwf')
        with fx.nav.FakeNav(doc):
            none_yet = reports.run('clash_delta', {}, doc=doc, interactive=False)
            entry, where = cr.record_snapshot([test_a, test_b], 'Before', doc)
            assert where == 'Kept on this computer.'
            r1 = test_a.Children[0]
            r1.Status = fx.nav._Enum('Resolved')                         # fixed by status
            test_b.Children.pop(0)                                       # Clash4 gone: fixed
            test_b.Children.append(fx.nav.ClashResult('Clash6', r1.Item2, r1.Item1, (4000.0, 0.0, 0.0)))
            test_b.Children[-1].Parent = test_b
            result = reports.run('clash_delta', {'folder': folder}, doc=doc, interactive=False)
            missing = reports.run('clash_delta', {'folder': folder, 'baseline': 'no such'}, doc=doc,
                                  interactive=False)
    assert none_yet['level'] == 'info' and none_yet['message'] == 'No clash snapshots yet'
    assert result['level'] == 'success', result
    assert result['message'] == '1 new, 2 fixed, 1 still open'
    assert result['detail'].startswith('Snapshot Before (') and 'Saved ' in result['detail']
    assert os.path.isfile(result['paths'][0])
    totals = result['delta']['totals']
    assert totals['then_open'] == 3 and totals['now_open'] == 2 and totals['changed'] == 1
    assert missing['level'] == 'warning' and missing['message'] == 'That snapshot is not there'


def test_unsaved_model():
    doc, _a, _b = fx.fake_clash_doc()
    with fx.Redirected('rep-clash-delta-unsaved'):
        with fx.nav.FakeNav(doc):
            result = reports.run('clash_delta', {}, doc=doc, interactive=False)
    assert result['message'] == 'Save the model first'
    assert datetime is not None
